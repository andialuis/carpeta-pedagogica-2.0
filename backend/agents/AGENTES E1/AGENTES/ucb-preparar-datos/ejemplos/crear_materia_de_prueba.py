"""Genera una materia sintetica de prueba (datos inventados, nadie real)."""
import csv, random, datetime as dt
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter

# Se crea junto a este archivo, nunca en la carpeta de trabajo del docente.
# Se puede pasar otro destino como primer argumento (lo usa instalar.py).
import sys
RAIZ = (Path(sys.argv[1]) if len(sys.argv) > 1
        else Path(__file__).resolve().parent / "materia_de_prueba")
M = RAIZ / "MATERIA SIS-12026"
M.mkdir(parents=True, exist_ok=True)
random.seed(7)

NOMBRES = [
    ("SIS-001", "Ana Lucia Rojas Camacho", "ana.rojas@ucb.edu.bo"),
    ("SIS-002", "Bruno Ernesto Paz Vargas", "bruno.paz@ucb.edu.bo"),
    ("SIS-003", "Carla Nunez Ortiz", "carla.nunez@ucb.edu.bo"),
    ("SIS-004", "Diego Alberto Mamani Quispe", "diego.mamani@ucb.edu.bo"),
    ("SIS-005", "Elena Sofia Terrazas Lima", "elena.terrazas@ucb.edu.bo"),
    ("SIS-006", "Fabio Ivan Cortez Salazar", "fabio.cortez@ucb.edu.bo"),
    ("SIS-007", "Gabriela Munoz Reyes", "gabriela.munoz@ucb.edu.bo"),
    ("SIS-008", "Hector Luis Aliaga Pinto", "hector.aliaga@ucb.edu.bo"),
]

# ---------- 1. Lista oficial de matricula (con columna N de fila, trampa) ----------
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Matricula"
ws.append(["N", "Codigo", "Apellidos y Nombres", "Correo institucional", "Carrera"])
for i, (cod, nom, correo) in enumerate(NOMBRES, 1):
    ws.append([i, cod, nom, correo, "Ingenieria de Sistemas"])
wb.save(M / "registro_matricula_oficial.xlsx")

# ---------- 2. Planilla de notas con encabezado combinado en 2 filas ----------
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Parciales"
ws["A1"] = "Estudiante"; ws.merge_cells("A1:A2")
ws["B1"] = "Primer Parcial"; ws.merge_cells("B1:C1")
ws["D1"] = "Segundo Parcial"; ws.merge_cells("D1:E1")
ws["B2"] = "Practica 1 s/15"; ws["C2"] = "Examen s/35"
ws["D2"] = "Practica 2 s/15"; ws["E2"] = "Examen s/35"
for cod, nom, correo in NOMBRES:
    ap, *resto = nom.split()
    invertido = f"{' '.join(nom.split()[-2:])}, {' '.join(nom.split()[:-2])}"  # apellido primero
    ws.append([invertido, random.randint(8,15), random.randint(18,35),
               random.randint(8,15), random.randint(18,35)])
ws2 = wb.create_sheet("Practicas")
ws2.append(["Correo", "Lab 1", "Lab 2", "Lab 3"])
for cod, nom, correo in NOMBRES[:6]:      # 2 alumnos ausentes a proposito
    ws2.append([correo, random.randint(5,10), random.randint(5,10), ""])
wb.save(M / "notas_parciales.xlsx")

# ---------- 3. Asistencia CSV, nombre abreviado ----------
fechas = ["2026-05-05","2026-05-12","2026-05-19","2026-05-26"]
with open(M / "asistencia_lunes.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["Nombre del estudiante"] + [f"Asistencia_{d}" for d in fechas])
    for cod, nom, correo in NOMBRES:
        partes = nom.split()
        abreviado = f"{partes[0]} {partes[1][0]}. {' '.join(partes[2:])}"
        w.writerow([abreviado] + [random.choice(["P","A","P","P"]) for _ in fechas])
    w.writerow(["Zoraida Fantasma Inexistente"] + ["P"]*4)   # fila huerfana a proposito

# ---------- 4. Registros del aula virtual (miles de filas) ----------
with open(M / "registros_logs.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["Hora", "Nombre completo", "Contexto del evento", "Componente"])
    base = dt.datetime(2026,3,1,8,0,0)
    for _ in range(1200):
        cod, nom, correo = random.choice(NOMBRES)
        t = base + dt.timedelta(days=random.randint(0,90), hours=random.randint(0,15))
        w.writerow([t.strftime("%d/%m/%Y %H:%M"), nom,
                    random.choice(["Curso: SIS-120","Foro","Tarea 1"]),
                    random.choice(["Sistema","Foro","Archivo","Cuestionario"])])

# ---------- 5. Silabo con ponderaciones (y un conflicto de escala) ----------
(M / "silabo_materia.txt").write_text(
    "SILABO SIS-120 PROGRAMACION\n\n"
    "SISTEMA DE EVALUACION\n"
    "Primer Parcial 50 puntos\n"
    "Segundo Parcial 50 puntos\n"
    "Practica 1 20 puntos\n"
    "Laboratorios 30 puntos\n"
    "Asistencia 10 %\n", encoding="utf-8")

# ---------- 6. Notas de defensa dentro de un documento ----------
(M / "acta_defensa_final.txt").write_text(
    "ACTA DE DEFENSA FINAL - calificaciones sobre 100\n\n"
    "1. Rojas Camacho, Ana Lucia - 88\n"
    "2. Paz Vargas, Bruno Ernesto - 74\n"
    "3. Nunez Ortiz, Carla - 91\n"
    "4. Mamani Quispe, Diego Alberto - 66\n"
    "5. Terrazas Lima, Elena Sofia - 82\n"
    "6. Cortez Salazar, Fabio Ivan - 70\n", encoding="utf-8")

# ---------- 7. Material NO CLAVE + imagen ----------
(M / "Tema_1_Introduccion.pptx").write_bytes(b"PK\x03\x04 fake pptx")
(M / "foto_notas_pizarra.png").write_bytes(b"\x89PNG\r\n\x1a\n fake")

# ---------- 8. Un acta en PDF ----------
def _pdf_minimo(lineas):
    """Un PDF de una pagina armado con biblioteca estandar.

    Comprueba que el programa enruta los PDF a transcripcion en vez de darlos por
    ilegibles. No hace falta ninguna libreria extra para generarlo.
    """
    texto = "BT /F1 12 Tf 72 720 Td 14 TL\n" + "\n".join(
        f"({l}) Tj T*" for l in lineas) + "\nET"
    objetos = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(texto)} >>\nstream\n{texto}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    salida, posiciones = "%PDF-1.4\n", []
    for numero, cuerpo in enumerate(objetos, start=1):
        posiciones.append(len(salida))
        salida += f"{numero} 0 obj\n{cuerpo}\nendobj\n"
    inicio_xref = len(salida)
    salida += f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n"
    for posicion in posiciones:
        salida += f"{posicion:010d} 00000 n \n"
    salida += (f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\n"
               f"startxref\n{inicio_xref}\n%%EOF\n")
    return salida.encode("latin-1")


(M / "acta_notas_firmada.pdf").write_bytes(_pdf_minimo([
    "ACTA DE CALIFICACIONES - SIS-120",
    "Primer Parcial 50 puntos",
    "Segundo Parcial 50 puntos",
]))
print("Materia de prueba creada en:", M)
print("\nEjecutala con:")
print(f'  python3 {Path(__file__).resolve().parent.parent}/scripts/organizar.py "MATERIA SIS-12026" --raiz "{RAIZ}"')
for p in sorted(M.iterdir()): print("  -", p.name)
