"""Skill: reportar

Escribe BASE_INTEGRADA.xlsx con sus hojas de auditoria, calcula la huella de
contenido y arma el informe y el resumen de control.
"""

import datetime as _dt
import hashlib

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ENCABEZADO = PatternFill("solid", fgColor="1F3864")
LETRA_ENCABEZADO = Font(color="FFFFFF", bold=True)


def _texto_valor(valor):
    if valor is None:
        return ""
    if isinstance(valor, _dt.datetime):
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(valor, _dt.date):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def huella_contenido(columnas_base, filas_base, indice_id):
    """Huella de los VALORES de la hoja BASE, ordenados por codigo de estudiante.

    No incluye los nombres de las columnas: dos programas pueden rotularlas
    distinto y tener exactamente los mismos datos.
    """
    lineas = []
    for fila in sorted(filas_base, key=lambda f: _texto_valor(f[indice_id])):
        lineas.append("".join(_texto_valor(v) for v in fila))
    crudo = "".join(lineas).encode("utf-8")
    return hashlib.sha256(crudo).hexdigest()[:16]


def _escribir_hoja(libro, titulo, columnas, filas):
    hoja = libro.create_sheet(titulo)
    hoja.append(list(columnas))
    for celda in hoja[1]:
        celda.fill = ENCABEZADO
        celda.font = LETRA_ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for fila in filas:
        hoja.append(list(fila))
    hoja.freeze_panes = "A2"
    for indice, columna in enumerate(columnas, start=1):
        ancho = max(12, min(46, len(str(columna)) + 2))
        hoja.column_dimensions[hoja.cell(row=1, column=indice).column_letter].width = ancho
    return hoja


def _plegar_columnas_de_identidad(hoja, columnas_base, diccionario):
    """Oculta las columnas de identidad repetidas de los archivos de origen.

    Cada archivo escribia al estudiante a su manera y R6 obliga a conservar todas
    esas columnas para poder auditar los cruces. Pero al abrir la base, ver cinco
    columnas de nombres desorienta. Se dejan visibles las de la lista oficial y se
    PLIEGAN las repetidas: siguen ahi, con todos sus datos, y se despliegan con el
    signo (+) de Excel. No se borra ni se altera ningun valor.
    """
    repetidas = {d["columna_en_base"] for d in diccionario
                 if d.get("papel") in ("identifica al estudiante",
                                       "metadato de procedencia")}
    if not repetidas:
        return 0

    plegadas = 0
    for indice, columna in enumerate(columnas_base, start=1):
        if columna not in repetidas:
            continue
        letra = hoja.cell(row=1, column=indice).column_letter
        dimension = hoja.column_dimensions[letra]
        dimension.hidden = True
        dimension.outlineLevel = 1
        plegadas += 1
    hoja.sheet_properties.outlinePr.summaryRight = False
    return plegadas


def contar_columnas_plegables(diccionario):
    """Cuantas columnas de identidad repetida se pliegan en la hoja BASE."""
    return len({d["columna_en_base"] for d in diccionario
                if d.get("papel") in ("identifica al estudiante",
                                      "metadato de procedencia")})


def _tabla(registros, campos):
    return [[r.get(c, "") for c in campos] for r in registros]


TITULO = Font(bold=True, size=13, color="1F3864")
SECCION = Font(bold=True, color="FFFFFF")


def _escribir_guia(libro, columnas_base, filas_base, diccionario, cobertura,
                   incidencias, reglas, fuente_oficial):
    """Hoja de lectura: que es este archivo y que hay que mirar primero.

    Una base con decenas de columnas es ilegible sin un mapa. Esta hoja no contiene
    ni un dato de estudiante: solo explica la estructura en lenguaje llano.
    """
    hoja = libro.create_sheet("GUIA_DE_LECTURA")
    hoja.column_dimensions["A"].width = 46
    hoja.column_dimensions["B"].width = 78

    def seccion(texto):
        fila = hoja.max_row + 2 if hoja.max_row > 1 else hoja.max_row + 1
        celda = hoja.cell(row=fila, column=1, value=texto)
        celda.font = SECCION
        celda.fill = ENCABEZADO
        hoja.cell(row=fila, column=2).fill = ENCABEZADO

    def linea(clave, valor=""):
        fila = hoja.max_row + 1
        hoja.cell(row=fila, column=1, value=clave)
        celda = hoja.cell(row=fila, column=2, value=valor)
        celda.alignment = Alignment(wrap_text=True, vertical="top")

    titulo = hoja.cell(row=1, column=1, value="COMO LEER ESTE ARCHIVO")
    titulo.font = TITULO

    seccion("LO ESENCIAL")
    linea("Una fila por estudiante", f"{len(filas_base)} estudiante(s) de la lista "
                                     f"oficial. Ninguna fila representa a dos "
                                     f"personas, ninguna persona ocupa dos filas.")
    linea("Lista oficial usada", fuente_oficial)
    linea("Columnas en total", str(len(columnas_base)))
    linea("Esta etapa NO limpia",
          "Los valores estan tal como venian: no se rellenaron vacios, no se "
          "convirtieron escalas, no se calculo ninguna nota. Eso es la Etapa 2.")

    seccion("QUE HAY EN CADA HOJA")
    for nombre, que in (
            ("BASE", "Los datos. Una fila por estudiante, todas las columnas."),
            ("DICCIONARIO", "De que archivo, hoja y columna salio cada dato."),
            ("COBERTURA", "En que situacion quedo cada estudiante en cada archivo."),
            ("EMPAREJAMIENTOS", "Con que metodo se reconocio a cada uno. Para auditar."),
            ("INCIDENCIAS", "Todo lo que no encajo, con la fila de origen completa."),
            ("REGLAS_EVALUACION", "Ponderaciones halladas y conflictos de escala.")):
        linea(nombre, que)

    seccion("COMO ESTAN ORGANIZADAS LAS COLUMNAS DE 'BASE'")
    linea("Primero, sin prefijo", "Las columnas de la lista oficial de matricula.")
    linea("Despues, con [prefijo]",
          "El prefijo dice de que archivo vino la columna. Si el archivo tenia "
          "varias hojas, aparece como [archivo · hoja].")
    linea("Columnas plegadas",
          "Cada archivo traia su propia columna de nombre o correo. Se conservan "
          "todas (hacen falta para auditar los cruces) pero estan PLEGADAS: se "
          "abren con el signo + arriba de las letras de columna en Excel.")

    bloques = {}
    for entrada in diccionario:
        archivo = entrada["archivo"]
        bloques.setdefault(archivo, {"total": 0, "identidad": 0, "datos": 0})
        bloques[archivo]["total"] += 1
        if entrada.get("papel") in ("identifica al estudiante",
                                    "metadato de procedencia"):
            bloques[archivo]["identidad"] += 1
        else:
            bloques[archivo]["datos"] += 1
    seccion("DE DONDE VIENE CADA BLOQUE DE COLUMNAS")
    for archivo, conteo in sorted(bloques.items()):
        linea(archivo, f"{conteo['datos']} columna(s) de datos"
                       + (f" + {conteo['identidad']} plegada(s)"
                          if conteo["identidad"] else ""))

    seccion("QUE REVISAR PRIMERO")
    prioridades = []
    conteo_inc = {}
    for registro in incidencias:
        conteo_inc[registro["tipo"]] = conteo_inc.get(registro["tipo"], 0) + 1
    urgentes = {
        "ambiguedad_resuelta_por_eliminacion":
            "Se adivino a quien pertenecia una fila. CONFIRMALO en INCIDENCIAS.",
        "colision_dos_filas_al_mismo_estudiante":
            "Dos filas cayeron sobre el mismo estudiante. La segunda esta en "
            "INCIDENCIAS sin integrarse: decide cual vale.",
        "fila_no_reconocida":
            "Hay filas con estudiantes que no estan en la lista oficial.",
        "fila_ambigua_sin_resolver":
            "Filas que podian ser de dos estudiantes distintos.",
        "formato_aprobado_que_no_se_puede_abrir":
            "Un archivo no se pudo leer: sus datos NO estan en la base.",
        "archivo_duplicado_identico":
            "Un archivo repetido no se integro dos veces. Solo para tu informacion.",
        "columna_calculada_sin_valor":
            "Una columna con formulas (Promedio, Total...) entro VACIA porque el "
            "archivo no traia el resultado guardado. Abre ese Excel, guardalo y "
            "vuelve a ejecutar.",
        "codigo_de_otro_sistema":
            "Un archivo numera a los estudiantes con codigos distintos a los de la "
            "matricula. Se cruzo por nombre, que es menos seguro: revisa "
            "EMPAREJAMIENTOS.",
        "columnas_con_valores_identicos":
            "Dos archivos distintos traen los mismos valores. Puede ser el mismo "
            "dato exportado dos veces: revisa si necesitas ambos.",
        "imagen_pendiente_de_transcripcion":
            "Una imagen quedo sin transcribir: sus datos NO estan en la base.",
    }
    for tipo, texto in urgentes.items():
        if conteo_inc.get(tipo):
            prioridades.append((f"{conteo_inc[tipo]} × {tipo}", texto))
    conflictos = [r for r in reglas if r.get("observacion")]
    if conflictos:
        prioridades.append((f"{len(conflictos)} conflicto(s) de evaluacion",
                            "La ponderacion declarada no coincide con la escala de "
                            "la planilla. Se reporta, NO se corrige. Ver "
                            "REGLAS_EVALUACION."))
    situaciones = {}
    for registro in cobertura:
        situaciones[registro["situacion"]] = situaciones.get(
            registro["situacion"], 0) + 1
    if situaciones.get("encontrado_vacio"):
        prioridades.append((f"{situaciones['encontrado_vacio']} × encontrado_vacio",
                            "El estudiante figura en el archivo pero sin ninguna "
                            "cifra. NO es lo mismo que estar ausente: la Etapa 2 "
                            "necesita esa diferencia."))
    if not prioridades:
        linea("Nada urgente", "No quedaron casos dudosos.")
    for clave, texto in prioridades:
        linea(clave, texto)

    hoja.freeze_panes = "A2"
    return hoja


def escribir_base(ruta, columnas_base, filas_base, diccionario, cobertura,
                  emparejamientos, incidencias, reglas, fuente_oficial=""):
    libro = openpyxl.Workbook()
    libro.remove(libro.active)

    _escribir_guia(libro, columnas_base, filas_base, diccionario, cobertura,
                   incidencias, reglas, fuente_oficial)

    hoja_base = _escribir_hoja(libro, "BASE", columnas_base, filas_base)
    _plegar_columnas_de_identidad(hoja_base, columnas_base, diccionario)

    _escribir_hoja(libro, "DICCIONARIO",
                   ["columna_en_base", "archivo", "hoja", "columna_original",
                    "escala_detectada", "papel"],
                   _tabla(diccionario, ["columna_en_base", "archivo", "hoja",
                                        "columna_original", "escala_detectada", "papel"]))

    _escribir_hoja(libro, "COBERTURA",
                   ["id_estudiante", "estudiante", "archivo", "situacion", "detalle"],
                   _tabla(cobertura, ["id_estudiante", "estudiante", "archivo",
                                      "situacion", "detalle"]))

    _escribir_hoja(libro, "EMPAREJAMIENTOS",
                   ["archivo", "fila_origen", "valor_original", "id_estudiante",
                    "estudiante", "metodo", "puntaje"],
                   _tabla(emparejamientos, ["archivo", "fila_origen", "valor_original",
                                            "id_estudiante", "estudiante", "metodo",
                                            "puntaje"]))

    _escribir_hoja(libro, "INCIDENCIAS",
                   ["tipo", "archivo", "fila_origen", "valor_original", "detalle",
                    "contenido_fila"],
                   _tabla(incidencias, ["tipo", "archivo", "fila_origen",
                                        "valor_original", "detalle", "contenido_fila"]))

    _escribir_hoja(libro, "REGLAS_EVALUACION",
                   ["instrumento", "valor_declarado", "unidad", "archivo", "linea",
                    "texto_original", "columnas_relacionadas", "escala_en_planilla",
                    "observacion"],
                   _tabla(reglas, ["instrumento", "valor_declarado", "unidad",
                                   "archivo", "linea", "texto_original",
                                   "columnas_relacionadas", "escala_en_planilla",
                                   "observacion"]))

    libro.save(ruta)
    return ruta


def escribir_informe(ruta, resumen):
    lineas = ["# Informe de preparacion de datos", ""]
    for titulo, contenido in resumen:
        lineas.append(f"## {titulo}")
        lineas.append("")
        if isinstance(contenido, list):
            lineas.extend(f"- {c}" for c in contenido)
        else:
            lineas.append(str(contenido))
        lineas.append("")
    ruta.write_text("\n".join(lineas), encoding="utf-8")
    return ruta
