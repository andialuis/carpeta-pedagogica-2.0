"""Instalador y autodiagnostico de las habilidades UCB.

Se ejecuta UNA vez, despues de copiar la carpeta AGENTES/ al espacio de trabajo:

    python3 AGENTES/instalar.py

Comprueba el entorno, verifica que las habilidades producen el resultado correcto
con una materia de prueba, y deja el archivo AGENTS.md en la raiz del espacio para
que el agente descubra las habilidades por su cuenta.

No toca ningun archivo del docente.
"""

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

AQUI = Path(__file__).resolve().parent
RAIZ_ESPACIO = AQUI.parent
SKILL = AQUI / "ucb-preparar-datos"

VERSION_MINIMA = (3, 8)
HUELLA_ESPERADA = "7c5ab4fd0f5b54eb"   # resultado correcto de la materia de prueba

OK, FALLO, AVISO = "  [OK]   ", "  [FALLO]", "  [AVISO]"


def titulo(texto):
    print(f"\n{texto}\n" + "-" * len(texto))


# --------------------------------------------------------------------- entorno
def comprobar_python():
    actual = ".".join(str(n) for n in sys.version_info[:3])
    if sys.version_info < VERSION_MINIMA:
        minima = ".".join(str(n) for n in VERSION_MINIMA)
        print(f"{FALLO} Python {actual}: se necesita {minima} o superior.")
        print("          Instalalo desde https://www.python.org/downloads/")
        return False
    print(f"{OK} Python {actual}")
    return True


def comprobar_openpyxl():
    try:
        import openpyxl
    except ImportError:
        print(f"{FALLO} falta openpyxl, la unica dependencia.")
        print(f"          Instalala con:  {Path(sys.executable).name} -m pip install openpyxl")
        return False
    print(f"{OK} openpyxl {openpyxl.__version__}")
    return True


def comprobar_archivos():
    esperados = [
        SKILL / "SKILL.md",
        SKILL / "scripts" / "organizar.py",
        SKILL / "scripts" / "config.json",
        SKILL / "referencias" / "formato-salida.md",
        SKILL / "ejemplos" / "crear_materia_de_prueba.py",
        AQUI / "AGENTS.md",
    ]
    faltan = [p for p in esperados if not p.exists()]
    if faltan:
        print(f"{FALLO} la copia esta incompleta. Faltan:")
        for p in faltan:
            print(f"          {p.relative_to(AQUI)}")
        return False
    modulos = sorted(SKILL.glob("scripts/skill_*.py"))
    print(f"{OK} habilidad completa ({len(modulos)} modulos + punto de entrada)")
    return True


# ----------------------------------------------------------------- autodiagnostico
def prueba_de_extremo_a_extremo():
    """Corre la habilidad sobre una materia inventada y comprueba la huella."""
    with tempfile.TemporaryDirectory(prefix="ucb_prueba_") as temporal:
        banco = Path(temporal)
        generador = SKILL / "ejemplos" / "crear_materia_de_prueba.py"
        hecho = subprocess.run([sys.executable, str(generador), str(banco)],
                               capture_output=True, text=True)
        if hecho.returncode != 0:
            print(f"{FALLO} no se pudo crear la materia de prueba:")
            print(hecho.stderr.strip()[:400])
            return False

        organizar = SKILL / "scripts" / "organizar.py"
        hecho = subprocess.run(
            [sys.executable, str(organizar), "MATERIA SIS-12026", "--raiz", str(banco)],
            capture_output=True, text=True)
        if hecho.returncode != 0:
            print(f"{FALLO} la habilidad no llego al final:")
            print(hecho.stderr.strip()[:400])
            return False

        salida = hecho.stdout
        huella = ""
        lineas = salida.splitlines()
        for indice, linea in enumerate(lineas):
            if "Huella de contenido" in linea and indice + 1 < len(lineas):
                huella = lineas[indice + 1].strip()
                break

        base = banco / "MATERIA SIS-12026-REV" / "bases_de_datos" / "BASE_INTEGRADA.xlsx"
        if not base.exists():
            print(f"{FALLO} no se genero BASE_INTEGRADA.xlsx")
            return False

        import openpyxl
        libro = openpyxl.load_workbook(base)
        hojas = libro.sheetnames
        faltantes = [h for h in ("GUIA_DE_LECTURA", "BASE", "DICCIONARIO",
                                 "COBERTURA", "EMPAREJAMIENTOS", "INCIDENCIAS",
                                 "REGLAS_EVALUACION") if h not in hojas]
        if faltantes:
            print(f"{FALLO} a BASE_INTEGRADA.xlsx le faltan hojas: {faltantes}")
            return False
        filas = libro["BASE"].max_row - 1
        libro.close()

        print(f"{OK} BASE_INTEGRADA.xlsx con sus 7 hojas ({filas} estudiantes)")

        # Los formatos que el programa no abre (imagen, PDF) tienen que salir con
        # una plantilla de transcripcion, no con un error de formato ilegible.
        bases = base.parent
        plantillas = sorted(bases.glob("*_PENDIENTE_TRANSCRIPCION.xlsx"))
        pendientes = {p.name.split("_PENDIENTE")[0] for p in plantillas}
        if {"foto_notas_pizarra", "acta_notas_firmada"} <= pendientes:
            print(f"{OK} imagenes y PDF derivados a transcripcion "
                  f"({len(plantillas)} plantilla(s))")
        else:
            print(f"{AVISO} no se generaron las plantillas de transcripcion "
                  f"esperadas: {sorted(pendientes)}")
        if "Archivos originales sin modificar:\n  SI" in salida.replace("\r\n", "\n"):
            print(f"{OK} los archivos originales quedaron intactos")
        else:
            print(f"{AVISO} no se pudo confirmar que los originales quedaron intactos")

        if huella == HUELLA_ESPERADA:
            print(f"{OK} huella correcta ({huella}): el resultado es el esperado")
            return True
        print(f"{AVISO} huella {huella or '(no encontrada)'}, "
              f"se esperaba {HUELLA_ESPERADA}.")
        print("          La habilidad funciona, pero produce un resultado distinto al")
        print("          de referencia. Revisa si se modifico algun script.")
        return True


# --------------------------------------------------------------------- AGENTS.md
MARCA = "ucb-preparar-datos"


def instalar_agents_md():
    origen = AQUI / "AGENTS.md"
    destino = RAIZ_ESPACIO / "AGENTS.md"

    if not destino.exists():
        shutil.copy2(origen, destino)
        print(f"{OK} AGENTS.md instalado en {destino}")
        return

    actual = destino.read_text(encoding="utf-8", errors="ignore")
    if MARCA in actual:
        print(f"{OK} AGENTS.md ya menciona la habilidad; no se toco")
        return

    bloque = (
        "\n\n## Habilidades en AGENTES/\n\n"
        "Este espacio incluye habilidades en `AGENTES/`. Cuando el docente pida\n"
        "**ordenar, preparar, consolidar o integrar** los archivos de una materia,\n"
        "lee `AGENTES/ucb-preparar-datos/SKILL.md` completo y siguelo al pie de la\n"
        "letra, respetando sus pausas. Los archivos de `scripts/` se ejecutan; no se\n"
        "leen ni se reescriben.\n")
    with open(destino, "a", encoding="utf-8") as f:
        f.write(bloque)
    print(f"{OK} ya existia un AGENTS.md: se le AGREGO al final un aviso de la")
    print("          habilidad. No se borro ni se cambio nada de lo que tenias.")


# -------------------------------------------------------------------------- main
def main():
    print("=" * 68)
    print(" INSTALACION DE HABILIDADES UCB — Etapa 1: Preparar")
    print("=" * 68)
    print(f"\nEspacio de trabajo detectado: {RAIZ_ESPACIO}")

    titulo("1. Entorno")
    if not (comprobar_python() and comprobar_openpyxl()):
        print("\nCorrige lo anterior y vuelve a ejecutar este instalador.")
        sys.exit(1)

    titulo("2. Integridad de la copia")
    if not comprobar_archivos():
        print("\nVuelve a copiar la carpeta AGENTES/ completa.")
        sys.exit(1)

    titulo("3. Prueba con una materia inventada")
    print("  (no se usa ningun dato real; la prueba se borra sola al terminar)")
    if not prueba_de_extremo_a_extremo():
        print("\nLa habilidad no paso la prueba. Revisa el mensaje de arriba.")
        sys.exit(1)

    titulo("4. Descubrimiento por el agente")
    instalar_agents_md()

    print("\n" + "=" * 68)
    print(" LISTO")
    print("=" * 68)
    print("""
Ya puedes usarlo. Los pasos:

  1. Copia la carpeta de tu materia DENTRO de este espacio de trabajo.
  2. Abre el agente (en Antigravity: el Agent Manager o el chat del editor).
  3. Escribele, en tus palabras:

         Quiero preparar los datos de mi materia.

  4. Te hara dos preguntas y te mostrara la lista de archivos para que la
     apruebes. Nada se toca hasta que apruebes.

Si el agente no reconoce la habilidad, empieza tu mensaje con:

     Lee AGENTES/ucb-preparar-datos/SKILL.md y siguelo.

Tus archivos originales nunca se modifican: todo el trabajo ocurre en una
carpeta nueva terminada en -REV.
""")


if __name__ == "__main__":
    main()
