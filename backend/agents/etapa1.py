"""
etapa1.py - Preparar y Consolidar Datos
Funciona con CUALQUIER materia. Recibe subject_name como argumento CLI.
Producto: BASE_INTEGRADA.xlsx + manifiesto.json
"""
import os, sys, json, re
from datetime import datetime

def run():
    print("Iniciando Etapa 1: Preparar y Consolidar Datos...")
    subject_name = sys.argv[1] if len(sys.argv) > 1 else None
    if not subject_name:
        print("Error: Se requiere el nombre de la materia como argumento.")
        sys.exit(1)

    agents_dir = os.path.dirname(__file__)
    backend_dir = os.path.dirname(agents_dir)
    root_dir = os.path.dirname(backend_dir)
    uploads_dir = os.path.join(root_dir, "uploads")
    raw_dir = os.path.join(uploads_dir, subject_name)
    rev_dir = os.path.join(uploads_dir, f"{subject_name}-REV")
    out_dir = os.path.join(rev_dir, "bases_de_datos")
    os.makedirs(out_dir, exist_ok=True)
    os.makedirs(os.path.join(rev_dir, "documentos"), exist_ok=True)
    os.makedirs(os.path.join(rev_dir, "imagenes"), exist_ok=True)
    os.makedirs(os.path.join(rev_dir, "presentaciones"), exist_ok=True)

    if not os.path.exists(raw_dir):
        print(f"Error: No se encontro la carpeta: {raw_dir}")
        sys.exit(1)

    print(f"Materia: {subject_name}")

    try:
        import pandas as pd
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError as e:
        print(f"Error importando librerias: {e}")
        sys.exit(1)

    excel_files = [f for f in os.listdir(raw_dir) if f.endswith(('.xlsx', '.xls'))]
    csv_files   = [f for f in os.listdir(raw_dir) if f.endswith('.csv')]
    print(f"Archivos: {len(excel_files)} Excel, {len(csv_files)} CSV")

    if not excel_files and not csv_files:
        print("Error: No hay archivos de datos.")
        sys.exit(1)

    student_col_candidates = ["nombre", "name", "alumno", "estudiante", "student"]
    def detect_student_col(df):
        for c in df.columns:
            if any(k in str(c).lower() for k in student_col_candidates):
                return c
        for c in df.columns:
            if df[c].dtype == object:
                return c
        return df.columns[0]

    dfs = {}
    for fname in excel_files:
        try:
            df = pd.read_excel(os.path.join(raw_dir, fname))
            if len(df) > 0:
                dfs[fname] = df
                print(f"  [OK] {fname} -> {len(df)} filas")
        except Exception as e:
            print(f"  [SKIP] {fname}: {e}")

    for fname in csv_files:
        try:
            df = pd.read_csv(os.path.join(raw_dir, fname), encoding='utf-8', errors='replace')
            if len(df) > 200:
                sc = detect_student_col(df)
                num_cols = df.select_dtypes(include='number').columns.tolist()
                if num_cols:
                    df_agg = df.groupby(sc)[num_cols].sum().reset_index()
                    df_agg.columns = [sc] + [f"Moodle_{c}" for c in num_cols]
                else:
                    df_agg = df.groupby(sc).size().reset_index(name='Total_Eventos_Moodle')
                dfs[fname] = df_agg
            else:
                dfs[fname] = df
            print(f"  [OK] {fname} -> {len(dfs[fname])} filas")
        except Exception as e:
            print(f"  [SKIP] {fname}: {e}")

    if not dfs:
        print("Error: Ningun archivo leido correctamente.")
        sys.exit(1)

    priority_keywords = ["lista", "matricula", "nomina", "estudiante", "alumno", "roster"]
    anchor_name = None
    for kw in priority_keywords:
        for fname in dfs:
            if kw in fname.lower():
                anchor_name = fname
                break
        if anchor_name:
            break
    if not anchor_name:
        anchor_name = max(dfs.keys(), key=lambda k: len(dfs[k]))

    base = dfs[anchor_name].copy()
    sc = detect_student_col(base)
    if sc != "Nombre":
        base = base.rename(columns={sc: "Nombre"})
    print(f"\nArchivo ancla: {anchor_name} ({len(base)} estudiantes)")

    for fname, df in dfs.items():
        if fname == anchor_name:
            continue
        df_copy = df.copy()
        s_col = detect_student_col(df_copy)
        if s_col != "Nombre":
            df_copy = df_copy.rename(columns={s_col: "Nombre"})
        existing = set(base.columns) - {"Nombre"}
        new_cols = [c for c in df_copy.columns if c != "Nombre" and c not in existing]
        if not new_cols:
            continue
        df_m = df_copy[["Nombre"] + new_cols]
        try:
            base = base.merge(df_m, on="Nombre", how="left")
            print(f"  [MERGE] {fname}: +{len(new_cols)} columnas")
        except Exception as e:
            print(f"  [SKIP merge] {fname}: {e}")

    print(f"\nBase integrada: {len(base)} estudiantes x {len(base.columns)} variables")

    cols_str = " ".join([str(c).lower() for c in base.columns])
    dimensions = {
        "rendimiento":               any(k in cols_str for k in ["nota","examen","quiz","parcial","promedio","proyecto","exposicion"]),
        "competencias_cbl":          any(k in cols_str for k in ["competencia","hito","cbl","nivel_global","rubrica","logro"]),
        "interaccion_digital":       any(k in cols_str for k in ["moodle","evento","click","asistencia","interaccion"]),
        "contexto_sociodemografico": any(k in cols_str for k in ["trabajo","conectividad","riesgo","distancia","socio","edad"]),
        "afectivo_motivacional":     any(k in cols_str for k in ["confianza","satisfaccion","sentimiento","autoevaluacion"]),
        "proceso_sudor_intelectual": any(k in cols_str for k in ["iteracion","prompt","autonomia","dialogo","copiloto","sudor"]),
        "diagnostico_inicial":       any(k in cols_str for k in ["diagnostico","nivel_inicial","pretest"]),
    }
    warnings = []
    if not dimensions["rendimiento"]:
        warnings.append("Sin datos de rendimiento: No se pueden calcular promedios.")
    if not dimensions["competencias_cbl"]:
        warnings.append("Sin datos CBL: Analisis de maestria competencial limitado.")
    if not dimensions["proceso_sudor_intelectual"]:
        warnings.append("Sin Sudor Intelectual: Riesgo de Outsourcing Cognitivo no detectable.")

    output_path = os.path.join(out_dir, "BASE_INTEGRADA.xlsx")
    base.to_excel(output_path, index=False)
    wb = openpyxl.load_workbook(output_path)
    ws = wb.active
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="1A3A5C")
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
    ws.freeze_panes = "B2"
    wb.save(output_path)

    manifest = {
        "subject": subject_name,
        "generated_at": datetime.now().isoformat(),
        "etapa": "1 - Preparar Datos",
        "total_students": len(base),
        "total_variables": len(base.columns),
        "dimensions": dimensions,
        "source_files": excel_files + csv_files,
        "warnings": warnings
    }
    with open(os.path.join(rev_dir, "manifiesto.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)

    # Inicializar y registrar en el Control de Versiones Documental (ISO 21001)
    try:
        sys.path.append(backend_dir)
        from version_control_service import init_version_control_workbook, record_document_version
        vc_file = init_version_control_workbook(subject_name)
        record_document_version(
            subject_name=subject_name,
            doc_name="BASE_INTEGRADA.xlsx",
            doc_code=f"DOC-CP2-BASE-{re.sub(r'[^a-zA-Z0-9]', '', subject_name)[:8]}",
            doc_type="Datos Curriculares y Analíticos",
            new_version="1.0",
            author_or_agent="Agente Etapa 1 - Preparador",
            change_description=f"Consolidación inicial de {len(excel_files) + len(csv_files)} archivos crudos para {len(base)} estudiantes.",
            justification="Aseguramiento de trazabilidad y limpieza de datos según norma ISO 21001.",
            file_path=output_path
        )
        record_document_version(
            subject_name=subject_name,
            doc_name="CONTROL_VERSIONES_DOCUMENTAL.xlsx",
            doc_code=f"DOC-CP2-VER-{re.sub(r'[^a-zA-Z0-9]', '', subject_name)[:8]}",
            doc_type="Registro Oficial de Control Documental",
            new_version="1.0",
            author_or_agent="Sistema de Gestión de Calidad",
            change_description="Apertura de la bitácora inmutable y matriz de versiones de la asignatura.",
            justification="Cumplimiento formal de la Cláusula 7.5 de ISO 21001 (Información Documentada).",
            file_path=vc_file
        )
        print(f"  [OK] Control de Versiones Documental inicializado: {vc_file}")
        print("  [AVISO CALIDAD] Si realiza cambios manuales en Word/Excel, actualice el número de versión en esta hoja.")
    except Exception as e:
        print(f"  [ADVERTENCIA] No se pudo actualizar control de versiones: {e}")

    for dim, present in dimensions.items():
        print(f"  {'[SI]' if present else '[NO]'} {dim}")
    if warnings:
        for w in warnings:
            print(f"  [!] {w}")

    print(f"\nEtapa 1 completada. Archivo: {output_path}")

if __name__ == "__main__":
    run()

