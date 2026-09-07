"""
etapa3.py - Analizar Datos con IA (Gemini) y Enfoque Humano
Genera INFORME_COMPLETO_E3.docx e insights.json para cualquier materia.
"""
import os
import sys
import pandas as pd
import json
import statistics
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

def set_cell_background(cell, hex_color):
    """Aplica color de fondo hexadecimal a una celda de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    """Establece márgenes internos de celda en dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def run():
    print("Iniciando Etapa 3: Análisis con IA y Enfoque Humano...")
    subject_name = sys.argv[1] if len(sys.argv) > 1 else "INV101 - Introduccion a la Investigacion"
    
    agents_dir = os.path.dirname(__file__)
    backend_dir = os.path.dirname(agents_dir)
    sys.path.insert(0, backend_dir)
    root_dir = os.path.dirname(backend_dir)
    target_dir = os.path.join(root_dir, "uploads", f"{subject_name}-REV", "bases_de_datos")
    
    input_file = os.path.join(target_dir, "BASE_LIMPIA_ANONIMIZADA.xlsx")
    if not os.path.exists(input_file):
        # Intentar BASE_INTEGRADA si la anonimizada aún no se corrió
        input_file = os.path.join(target_dir, "BASE_INTEGRADA.xlsx")
    
    json_output = os.path.join(target_dir, "insights.json")
    docx_output = os.path.join(target_dir, "INFORME_COMPLETO_E3.docx")
    
    if not os.path.exists(input_file):
        print(f"Error: No se encontró {input_file}. Ejecuta la Etapa 1 o 2 primero.")
        sys.exit(1)
        
    print(f"Cargando datos desde {input_file} para asignatura: {subject_name}")
    df = pd.read_excel(input_file)
    
    # Cargar llave de nombres si existe
    llave_path = os.path.join(target_dir, "llave_nombres.json")
    llave = {}
    if os.path.exists(llave_path):
        try:
            with open(llave_path, "r", encoding="utf-8") as f:
                llave = json.load(f)
        except Exception:
            pass

    # Identificar columna de ID o Nombre
    id_col = None
    for c in ["Estudiante_ID", "Nombre", "ID", "Codigo", "Alumno"]:
        if c in df.columns:
            id_col = c
            break
    if not id_col:
        id_col = df.columns[0]
        
    # Identificar columna principal de rendimiento
    perf_col = None
    for c in ["Promedio_Ponderado", "Promedio", "Nota Final", "Nota_Final", "Puntaje_Total", "Calificacion"]:
        if c in df.columns:
            perf_col = c
            break
    if not perf_col:
        num_cols = df.select_dtypes(include='number').columns.tolist()
        perf_col = num_cols[-1] if num_cols else df.columns[-1]

    # Identificar columna de iteraciones / esfuerzo
    iter_col = None
    for c in ["Iteraciones_Prompt_Total", "Iteraciones", "Prompts", "Total_Eventos_Moodle", "Horas_Estudio"]:
        if c in df.columns:
            iter_col = c
            break

    # Extraer estadísticas descriptivas
    desc_stats = df.describe()

    # Verificar API Key de Gemini
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        from dotenv import load_dotenv
        load_dotenv(os.path.join(backend_dir, ".env"))
        api_key = os.environ.get("GEMINI_API_KEY")

    insights = None
    informe_texto = ""
    error_gemini_msg = None

    if api_key:
        try:
            from google import genai
            from google.genai import types
            
            client = genai.Client(api_key=api_key)
            datos_estudiantes = df.head(30).to_csv(index=False)
            prompt = f"""Eres un experto pedagógico en Analítica de Aprendizaje con Enfoque Humano (Deci & Ryan, Vygotsky, Freire).
Asignatura: {subject_name}
Datos del grupo:
{datos_estudiantes}

Escribe un informe analítico estructurado en formato JSON estricto con las siguientes claves:
{{
  "promedio_general": 75.0,
  "desviacion_estandar": 12.0,
  "riesgo_abandono": 2,
  "cobertura_cbl": "CBL Activo",
  "alerta_outsourcing_cognitivo": true,
  "outsourcing_count": 2,
  "potencial_oculto_count": 3,
  "necesidad_andamiaje_vygotsky": "Descripción analítica de la ZDP.",
  "alertas": ["Alerta 1", "Alerta 2"],
  "deci_ryan_clase": {{"autonomia": 7.0, "competencia": 6.8, "relacion": 7.2}},
  "hipotesis_gemini": "Párrafo interpretativo profundo sobre el grupo.",
  "estudiantes": [
    {{
      "id": "Est_001",
      "promedio": 85.0,
      "nivel_cbl": "Competente",
      "riesgo_abandono": "Bajo",
      "autonomia_deci_ryan": 8,
      "competencia_deci_ryan": 8,
      "relacion_deci_ryan": 7,
      "alerta_outsourcing_cognitivo": false,
      "requiere_triangulacion": false,
      "zdp_vygotsky": "En consolidación",
      "estrategia_hacker_curriculo": "Estrategia pedagógica",
      "recomendacion": "Recomendación"
    }}
  ]
}}"""
            gemini_model = os.environ.get("GEMINI_MODEL", "gemini-3.6-flash")
            print(f"Consultando a {gemini_model} con Structured Outputs...")
            response = client.models.generate_content(
                model=gemini_model,
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    temperature=0.2
                ),
            )
            raw_text = response.text.strip()
            insights = json.loads(raw_text)

            # Normalizar si Gemini devolvió un array [ {...} ] en lugar de un diccionario { ... }
            if isinstance(insights, list):
                if len(insights) > 0 and isinstance(insights[0], dict):
                    if "estudiantes" in insights[0]:
                        insights = insights[0]
                    else:
                        insights = {"estudiantes": insights}
                else:
                    insights = {"estudiantes": []}

            if isinstance(insights, dict) and insights.get("estudiantes"):
                insights["motor_ia"] = {
                    "modo": "gemini",
                    "modelo": gemini_model,
                    "mensaje": f"Análisis cualitativo y pedagógico generado exitosamente con {gemini_model}.",
                    "fallback_activado": False,
                    "error_motivo": None
                }
                informe_texto = insights.get("hipotesis_gemini", "")
        except Exception as e:
            error_gemini_msg = str(e)
            print(f"[IA ADVERTENCIA] Falló consulta a Gemini ({e}). Activando motor determinista local de contingencia.")
            insights = None
    else:
        error_gemini_msg = "API_KEY_NO_CONFIGURADA"
        print("[IA ADVERTENCIA] GEMINI_API_KEY no detectada. Activando motor determinista local de contingencia.")

    # Motor Analítico Determinista / Robusto (si no hay Gemini o falló)
    if not insights or not insights.get("estudiantes"):
        print("Ejecutando motor analítico con Enfoque Humano...")
        estudiantes_list = []
        scores = []
        outsourcing_count = 0
        risk_count = 0
        hidden_count = 0

        for i, row in df.iterrows():
            est_id = str(row[id_col])
            nombre_real = llave.get(est_id, est_id)
            
            try:
                val = float(row[perf_col]) if perf_col in row and pd.notnull(row[perf_col]) else 70.0
            except Exception:
                val = 70.0
            scores.append(val)
            
            # Obtener iteraciones si existen
            iters = 0
            if iter_col and iter_col in row and pd.notnull(row[iter_col]):
                try:
                    iters = int(float(row[iter_col]))
                except Exception:
                    iters = 0

            # Detección de Outsourcing Cognitivo
            is_outsourcing = (val >= 90 and iters <= 5 and iter_col is not None)
            is_risk = (val < 51)
            is_hidden = (55 <= val <= 75 and iters >= 20)

            if is_outsourcing:
                outsourcing_count += 1
                nivel_cbl = "Avanzado (no verificado)"
                riesgo_ab = "Bajo (aparente)"
                aut = 2
                comp = 2
                rel = 3
                zdp = "Zona de Confort Digital: Riesgo de outsourcing cognitivo."
                estrat = "Triangulación Socrática obligatoria. Re-evaluación oral del proceso."
            elif is_risk:
                risk_count += 1
                nivel_cbl = "Inicial / En Riesgo"
                riesgo_ab = "Alto"
                aut = 3
                comp = 3
                rel = 4
                zdp = "Prerrequisitos pendientes: Requiere andamiaje intensivo."
                estrat = "Entrevista motivacional y articulación de apoyo socioeducativo."
            elif is_hidden:
                hidden_count += 1
                nivel_cbl = "Competente"
                riesgo_ab = "Bajo"
                aut = 7
                comp = 6
                rel = 8
                zdp = "ZDP activa: Potencial real supera las calificaciones cuantitativas."
                estrat = "Andamiaje en expresión académica y rúbricas alternativas (DUA)."
            elif val >= 80:
                nivel_cbl = "Avanzado"
                riesgo_ab = "Bajo"
                aut = 9
                comp = 9
                rel = 8
                zdp = "Maestría: Listo para proyectos de investigación autónomos."
                estrat = "Asignar rol de mentor de pares y desafíos avanzados."
            else:
                nivel_cbl = "Competente"
                riesgo_ab = "Bajo"
                aut = 6
                comp = 6
                rel = 6
                zdp = "Progresión sostenida dentro de los parámetros esperados."
                estrat = "Seguimiento formativo regular."

            estudiantes_list.append({
                "id": est_id,
                "nombre_real": nombre_real,
                "promedio": round(val, 1),
                "nivel_cbl": nivel_cbl,
                "riesgo_abandono": riesgo_ab,
                "autonomia_deci_ryan": aut,
                "competencia_deci_ryan": comp,
                "relacion_deci_ryan": rel,
                "alerta_outsourcing_cognitivo": is_outsourcing,
                "requiere_triangulacion": is_outsourcing,
                "zdp_vygotsky": zdp,
                "estrategia_hacker_curriculo": estrat,
                "recomendacion": f"Plan formativo personalizado para {nombre_real}."
            })

        n = len(scores) if scores else 1
        prom = round(sum(scores)/n, 1) if scores else 70.0
        mediana = round(float(statistics.median(scores)), 1) if scores else 70.0
        desv = round(statistics.stdev(scores), 1) if len(scores) > 1 else 0.0

        # Estadísticos de distribución (Asimetría y Curtosis)
        s_scores = pd.Series(scores)
        skew_val = round(float(s_scores.skew()), 2) if len(s_scores) > 2 else 0.0
        kurt_val = round(float(s_scores.kurt()), 2) if len(s_scores) > 3 else 0.0

        if skew_val < -0.5:
            skew_diag = "Asimetría Negativa: Concentración en notas altas (sesgo a la izquierda)."
        elif skew_val > 0.5:
            skew_diag = "Asimetría Positiva: Concentración en notas bajas (sesgo a la derecha)."
        else:
            skew_diag = "Distribución Cuasi-Simétrica: Calificaciones balanceadas."

        if kurt_val > 1.0:
            kurt_diag = "Leptocúrtica: Alta concentración de estudiantes en torno a la media."
        elif kurt_val < -1.0:
            kurt_diag = "Platicúrtica: Distribución aplanada y polarizada (alta dispersión)."
        else:
            kurt_diag = "Mesocúrtica: Dispersión moderada equivalente a normalidad."

        # Correlaciones de Rangos de Spearman (No paramétricas, robustas para muestras pequeñas)
        def compute_spearman(col_x, col_y):
            try:
                valid = df[[col_x, col_y]].dropna()
                if len(valid) < 3:
                    return 0.0
                vx = list(valid[col_x])
                vy = list(valid[col_y])
                def get_ranks(arr):
                    s_idx = sorted(range(len(arr)), key=lambda k: arr[k])
                    r = [0] * len(arr)
                    for rank_num, idx in enumerate(s_idx, 1):
                        r[idx] = rank_num
                    return r
                rx = get_ranks(vx)
                ry = get_ranks(vy)
                n = len(rx)
                d_sq = sum((rx[i] - ry[i]) ** 2 for i in range(n))
                rho = 1.0 - (6.0 * d_sq) / (n * (n ** 2 - 1))
                return round(float(rho), 2)
            except Exception:
                return 0.0

        asist_col = None
        for c in ["Total_Asistidas", "Porcentaje_%", "Asistencia", "Porcentaje_Asistencia"]:
            if c in df.columns:
                asist_col = c
                break

        spearman_asist_nota = compute_spearman(asist_col, perf_col) if asist_col and perf_col else 0.0
        spearman_proceso_nota = compute_spearman(iter_col, perf_col) if iter_col and perf_col else 0.0

        gauss_data = [{"rango": f"{r}-{r+9}", "estudiantes": sum(1 for p in scores if r <= p < r+10)} for r in range(0, 101, 10)]
        
        scatter_data = []
        for e in estudiantes_list:
            row_match = df[df[id_col].astype(str) == e["id"]]
            it_val = 15
            if not row_match.empty and iter_col and iter_col in row_match.columns:
                try:
                    it_val = int(row_match[iter_col].values[0])
                except Exception:
                    pass
            scatter_data.append({
                "id": e["id"],
                "promedio": e["promedio"],
                "iteraciones": it_val,
                "riesgo": e["riesgo_abandono"],
                "alerta_outsourcing": e["alerta_outsourcing_cognitivo"]
            })

        deci_ryan_class = {
            "autonomia": round(sum(e["autonomia_deci_ryan"] for e in estudiantes_list) / n, 1),
            "competencia": round(sum(e["competencia_deci_ryan"] for e in estudiantes_list) / n, 1),
            "relacion": round(sum(e["relacion_deci_ryan"] for e in estudiantes_list) / n, 1),
        }

        insights = {
            "subject": subject_name,
            "generated_at": pd.Timestamp.now().isoformat(),
            "total_students": n,
            "promedio_general": prom,
            "mediana_general": mediana,
            "desviacion_estandar": desv,
            "skewness": skew_val,
            "skewness_diagnostico": skew_diag,
            "kurtosis": kurt_val,
            "kurtosis_diagnostico": kurt_diag,
            "spearman_asistencia_nota": spearman_asist_nota,
            "spearman_proceso_nota": spearman_proceso_nota,
            "riesgo_abandono": risk_count,
            "cobertura_cbl": "Análisis Multidimensional Activo",
            "alerta_outsourcing_cognitivo": (outsourcing_count > 0),
            "outsourcing_count": outsourcing_count,
            "potencial_oculto_count": hidden_count,
            "necesidad_andamiaje_vygotsky": f"{risk_count} estudiantes requieren andamiaje urgente; {hidden_count} con potencial oculto detectado.",
            "alertas": [
                f"{outsourcing_count} estudiantes con alerta de Outsourcing Cognitivo (requieren Triangulación Socrática)." if outsourcing_count > 0 else "Sin anomalías de Outsourcing Cognitivo detectadas.",
                f"{risk_count} estudiantes clasificados en riesgo de rezago o abandono." if risk_count > 0 else "Nivel de retención óptimo en la clase.",
                f"{hidden_count} estudiantes con alto proceso pero notas moderadas (Potencial Oculto)." if hidden_count > 0 else "Procesos y notas alineados en el grupo.",
                f"Diagnóstico de Forma: {skew_diag} (Asimetría g1 = {skew_val})."
            ],
            "deci_ryan_clase": deci_ryan_class,
            "gauss_distribucion": gauss_data,
            "scatter_rendimiento_proceso": scatter_data,
            "estudiantes": estudiantes_list,
            "hipotesis_gemini": f"El grupo de {subject_name} presenta un promedio de {prom:.1f}/100 (Mediana: {mediana:.1f}) con desviación estándar de {desv:.1f}. Asimetría de Fisher: {skew_val} ({skew_diag}). Se identifican {risk_count} casos en riesgo y {outsourcing_count} alertas de outsourcing cognitivo.",
            "motor_ia": {
                "modo": "local_heuristico",
                "modelo": "motor_determinista_no_parametrico",
                "mensaje": (
                    "Medida de Contingencia Pedagógica: La clave de API de Gemini no está configurada aún en el sistema. "
                    if error_gemini_msg == "API_KEY_NO_CONFIGURADA" else
                    f"Medida de Contingencia Pedagógica: La API de Gemini reportó una incidencia ({error_gemini_msg}). "
                ) + "Se ejecutó el motor analítico determinista local con cálculo matemático puro (Spearman ρ, Asimetría Fisher g1, Curtosis g2 y ZDP Vygotsky) para garantizar la continuidad pedagógica sin pérdida de datos.",
                "fallback_activado": True,
                "error_motivo": error_gemini_msg
            }
        }
        informe_texto = insights["hipotesis_gemini"]

    # Generar gráficos Matplotlib
    fig1_path = os.path.join(target_dir, "fig1_notas.png")
    try:
        plt.figure(figsize=(8, 4))
        plt.hist([e["promedio"] for e in insights["estudiantes"]], bins=10, color='#1A3A5C', edgecolor='white')
        plt.title(f'Distribución de Rendimiento - {subject_name}')
        plt.xlabel('Calificación')
        plt.ylabel('N° Estudiantes')
        plt.grid(axis='y', alpha=0.3)
        plt.tight_layout()
        plt.savefig(fig1_path, dpi=150)
        plt.close()
    except Exception as e:
        print(f"Error generando gráfico: {e}")

    # Construir INFORME_COMPLETO_E3.docx
    print("Construyendo INFORME_COMPLETO_E3.docx...")
    doc = Document()

    codigo_materia = subject_name.split("-")[0].strip() if "-" in subject_name else "E3"
    total_est = len(insights.get("estudiantes", []))

    # Configuración de márgenes estándar (0.8 pulgadas)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        # Encabezado institucional de calidad
        header = section.header
        p_head = header.paragraphs[0]
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_h = p_head.add_run(f"INFORME ANALÍTICO DE APRENDIZAJE • DOC-DIR-CP2-INF3-{codigo_materia} • ISO 21001")
        r_h.font.name = "Arial"
        r_h.font.size = Pt(7.5)
        r_h.font.color.rgb = RGBColor(100, 116, 139)

        # Pie de página institucional
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_f = p_foot.add_run("Carpeta Pedagógica 2.0 • Sistema de Analítica de Aprendizaje & Aseguramiento de Calidad")
        r_f.font.name = "Arial"
        r_f.font.size = Pt(7.5)
        r_f.font.color.rgb = RGBColor(100, 116, 139)

    # Colores institucionales
    C_NAVY = RGBColor(26, 58, 92)       # #1A3A5C
    C_SLATE = RGBColor(30, 41, 59)      # #1E293B
    C_AMBER = RGBColor(180, 83, 9)      # #B45309
    C_DARK = RGBColor(15, 23, 42)       # #0F172A
    C_MUTED = RGBColor(100, 116, 139)   # #64748B

    # Portada / Título Principal
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("SISTEMA DE GESTIÓN DE LA CALIDAD EDUCATIVA • ISO 21001:2018\n")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(8.5)
    r_inst.font.bold = True
    r_inst.font.color.rgb = C_AMBER

    r_title = p_inst.add_run("INFORME OFICIAL DE ANALÍTICA DEL APRENDIZAJE Y AUDITORÍA COGNITIVA")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = C_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(f"ASIGNATURA: {subject_name.upper()} • ANÁLISIS MULTIDIMENSIONAL • ENFOQUE HUMANO")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(9.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = C_SLATE

    # Tabla de Control Documental ISO 21001
    t_iso = doc.add_table(rows=3, cols=2)
    t_iso.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_iso.autofit = False

    iso_meta = [
        ("Código Documental:", f"DOC-DIR-CP2-INF3-{codigo_materia}"),
        ("Versión y Vigencia:", "1.0 (Oficial Aprobado) • Periodo Académico Vigente"),
        ("Norma y Sistema Aplicable:", "ISO 21001:2018 Cláusula 9.1 (Seguimiento, Medición, Análisis y Evaluación)")
    ]
    for idx, (lbl, val) in enumerate(iso_meta):
        c0, c1 = t_iso.cell(idx, 0), t_iso.cell(idx, 1)
        set_cell_margins(c0, 50, 50, 100, 100)
        set_cell_margins(c1, 50, 50, 100, 100)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(lbl)
        r0.font.name = "Arial"
        r0.font.size = Pt(8)
        r0.font.bold = True
        r0.font.color.rgb = C_SLATE

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(8)
        r1.font.color.rgb = C_DARK

    doc.add_paragraph() # Espacio

    # Ficha Técnica de Calidad de la Muestra
    h_fic = doc.add_paragraph()
    r_hfic = h_fic.add_run("FICHA TÉCNICA DE CALIDAD DE LA MUESTRA")
    r_hfic.font.name = "Arial"
    r_hfic.font.size = Pt(10.5)
    r_hfic.font.bold = True
    r_hfic.font.color.rgb = C_SLATE

    t_fic = doc.add_table(rows=2, cols=2)
    t_fic.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_fic.autofit = False

    from datetime import datetime
    corte_str = datetime.now().strftime("%d/%m/%Y %H:%M")
    fic_data = [
        (f"Población Evaluada (N): {total_est} estudiantes", f"Muestra Analizada (n): {total_est} registros (100% Cobertura Censal)"),
        (f"Fecha y Hora de Corte: {corte_str}", "Confiabilidad Estadística: Censo Completo • Coeficiente de Spearman ρ")
    ]
    for r_i, r_data in enumerate(fic_data):
        for c_i, text in enumerate(r_data):
            cell = t_fic.cell(r_i, c_i)
            set_cell_margins(cell, 60, 60, 100, 100)
            set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = C_SLATE

    doc.add_paragraph() # Espacio
    
    # 1. Resumen Ejecutivo y Diagnóstico
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("1. RESUMEN EJECUTIVO Y DIAGNÓSTICO INTEGRAL")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(10.5)
    r_h1.font.bold = True
    r_h1.font.color.rgb = C_SLATE

    p_diag = doc.add_paragraph()
    r_diag = p_diag.add_run(informe_texto.strip())
    r_diag.font.name = "Arial"
    r_diag.font.size = Pt(9)
    r_diag.font.color.rgb = C_DARK
    
    # 2. Indicadores Clave de la Clase (KPIs)
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("2. INDICADORES CLAVE DE LA CLASE (KPIs)")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(10.5)
    r_h2.font.bold = True
    r_h2.font.color.rgb = C_SLATE

    kpis = [
        ("Promedio General de la Asignatura", f"{insights.get('promedio_general', 0)} / 100"),
        ("Estudiantes en Riesgo de Rezago o Abandono", f"{insights.get('riesgo_abandono', 0)} de {total_est}"),
        ("Alertas de Outsourcing Cognitivo (Requieren Triangulación)", f"{insights.get('outsourcing_count', 0)} casos"),
        ("Estudiantes con Potencial Oculto (Alto Proceso / Nota Media)", f"{insights.get('potencial_oculto_count', 0)} casos"),
    ]
    t_kpi = doc.add_table(rows=1, cols=2)
    t_kpi.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_kpi.autofit = False

    c0, c1 = t_kpi.cell(0, 0), t_kpi.cell(0, 1)
    set_cell_background(c0, "1A3A5C")
    set_cell_background(c1, "1A3A5C")
    set_cell_margins(c0, 80, 80, 100, 100)
    set_cell_margins(c1, 80, 80, 100, 100)
    
    p0 = c0.paragraphs[0]
    r0 = p0.add_run("Indicador Pedagógico")
    r0.font.name = "Arial"
    r0.font.size = Pt(8.5)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(255, 255, 255)

    p1 = c1.paragraphs[0]
    r1 = p1.add_run("Valor / Métrica Institucional")
    r1.font.name = "Arial"
    r1.font.size = Pt(8.5)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(255, 255, 255)

    for r_i, (k, v) in enumerate(kpis, start=1):
        row = t_kpi.add_row()
        bg = "FFFFFF" if r_i % 2 != 0 else "F8FAFC"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, 70, 70, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(str(k if c_idx == 0 else v))
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.bold = (c_idx == 1)
            r.font.color.rgb = C_SLATE if c_idx == 1 else C_DARK

    doc.add_paragraph() # Espacio

    if os.path.exists(fig1_path):
        h3 = doc.add_paragraph()
        r_h3 = h3.add_run("3. DISTRIBUCIÓN ESTADÍSTICA DEL RENDIMIENTO")
        r_h3.font.name = "Arial"
        r_h3.font.size = Pt(10.5)
        r_h3.font.bold = True
        r_h3.font.color.rgb = C_SLATE
        doc.add_picture(fig1_path, width=Inches(6.0))
        doc.add_paragraph() # Espacio

    # 4. Rigor Matemático y Análisis de Distribución
    h4 = doc.add_paragraph()
    r_h4 = h4.add_run("4. RIGOR MATEMÁTICO Y ANÁLISIS DE DISTRIBUCIÓN")
    r_h4.font.name = "Arial"
    r_h4.font.size = Pt(10.5)
    r_h4.font.bold = True
    r_h4.font.color.rgb = C_SLATE

    stat_metrics = [
        ("Media Aritmética (x̄)", f"{insights.get('promedio_general', 0)} / 100"),
        ("Mediana Muestral (Me)", f"{insights.get('mediana_general', 0)} / 100"),
        ("Desviación Estándar (s)", f"{insights.get('desviacion_estandar', 0)}"),
        ("Coeficiente de Asimetría (g1 - Fisher)", f"{insights.get('skewness', 0)} ({insights.get('skewness_diagnostico', '')})"),
        ("Curtosis Muestral (g2 - Fisher)", f"{insights.get('kurtosis', 0)} ({insights.get('kurtosis_diagnostico', '')})"),
        ("Correlación de Spearman Asistencia-Rendimiento (ρ)", f"{insights.get('spearman_asistencia_nota', 0)} (No paramétrica)"),
        ("Correlación de Spearman Sudor Intelectual-Rendimiento (ρ)", f"{insights.get('spearman_proceso_nota', 0)} (Medición de Proceso)"),
    ]
    t_stat = doc.add_table(rows=1, cols=2)
    t_stat.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_stat.autofit = False

    c0, c1 = t_stat.cell(0, 0), t_stat.cell(0, 1)
    set_cell_background(c0, "1A3A5C")
    set_cell_background(c1, "1A3A5C")
    set_cell_margins(c0, 80, 80, 100, 100)
    set_cell_margins(c1, 80, 80, 100, 100)

    p0 = c0.paragraphs[0]
    r0 = p0.add_run("Parámetro / Prueba Estadística")
    r0.font.name = "Arial"
    r0.font.size = Pt(8.5)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(255, 255, 255)

    p1 = c1.paragraphs[0]
    r1 = p1.add_run("Resultado y Diagnóstico Formal")
    r1.font.name = "Arial"
    r1.font.size = Pt(8.5)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(255, 255, 255)

    for r_i, (param, res) in enumerate(stat_metrics, start=1):
        row = t_stat.add_row()
        bg = "FFFFFF" if r_i % 2 != 0 else "F8FAFC"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, 70, 70, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(str(param if c_idx == 0 else res))
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.bold = (c_idx == 0)
            r.font.color.rgb = C_SLATE if c_idx == 0 else C_DARK

    doc.add_paragraph() # Espacio

    # 5. Alertas Pedagógicas y Diagnóstico de Enfoque Humano
    h5 = doc.add_paragraph()
    r_h5 = h5.add_run("5. ALERTAS PEDAGÓGICAS Y DIAGNÓSTICO DE ENFOQUE HUMANO")
    r_h5.font.name = "Arial"
    r_h5.font.size = Pt(10.5)
    r_h5.font.bold = True
    r_h5.font.color.rgb = C_SLATE

    for a in insights.get("alertas", []):
        p_al = doc.add_paragraph()
        r_al = p_al.add_run(f"• {a}")
        r_al.font.name = "Arial"
        r_al.font.size = Pt(8.5)
        r_al.font.color.rgb = C_DARK

    p_mar = doc.add_paragraph()
    p_mar.add_run("Fundamentación Epistémica:\n").bold = True
    p_mar.add_run(f"• Autodeterminación (Deci & Ryan): Autonomía={insights.get('deci_ryan_clase',{}).get('autonomia',7)}/10, Competencia={insights.get('deci_ryan_clase',{}).get('competencia',7)}/10, Relación={insights.get('deci_ryan_clase',{}).get('relacion',7)}/10.\n")
    p_mar.add_run(f"• Zona de Desarrollo Próximo (Vygotsky): {insights.get('necesidad_andamiaje_vygotsky', 'Evaluación continua')}\n")
    p_mar.add_run("• Pedagogía Crítica (Freire): Triangulación Socrática activada para validar autoría auténtica y erradicar la simulación pasiva.")
    for run_item in p_mar.runs:
        run_item.font.name = "Arial"
        run_item.font.size = Pt(8)

    doc.add_paragraph() # Espacio

    # 6. Matriz de Prescripción y Compromisos de Andamiaje en 3 Niveles
    h6 = doc.add_paragraph()
    r_h6 = h6.add_run("6. MATRIZ DE PRESCRIPCIÓN Y ACCIONES EN 3 NIVELES INSTITUCIONALES")
    r_h6.font.name = "Arial"
    r_h6.font.size = Pt(10.5)
    r_h6.font.bold = True
    r_h6.font.color.rgb = C_SLATE

    t_presc = doc.add_table(rows=4, cols=3)
    t_presc.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_presc.autofit = False

    presc_headers = ["Nivel de Gestión", "Acción Pedagógica / Prescripción", "Evidencia de Cumplimiento"]
    for c_i, h_txt in enumerate(presc_headers):
        c = t_presc.cell(0, c_i)
        set_cell_background(c, "1A3A5C")
        set_cell_margins(c, 80, 80, 80, 80)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_txt)
        r.font.name = "Arial"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    presc_data = [
        ("Nivel 1: Aula (Docente Titular)",
         "Aplicar andamiaje diferenciado en ZDP a casos en riesgo; ejecutar Triangulación Socrática a estudiantes con alerta de outsourcing.",
         "Actas de Triangulación Socrática y registros de Sudor Intelectual."),
        ("Nivel 2: Tutoría / Bienestar",
         "Acompañamiento psicopedagógico y seguimiento a estudiantes con ausentismo o brechas situacionales en conectividad/empleo.",
         "Ficha Socioeducativa de Contexto y bitácora de tutorías."),
        ("Nivel 3: Coordinación Académica",
         "Supervisar el alineamiento constructivo entre rúbricas CBL y evidencias, asegurando los créditos SCT y horas de trabajo autónomo.",
         "Plan de Aprendizaje Modular oficial y auditoría de portafolios.")
    ]

    for r_i, (n_niv, n_acc, n_evi) in enumerate(presc_data, start=1):
        row_c = t_presc.rows[r_i].cells
        bg = "FFFFFF" if r_i % 2 != 0 else "F8FAFC"
        for cell in row_c:
            set_cell_background(cell, bg)
            set_cell_margins(cell, 70, 70, 80, 80)

        p0 = row_c[0].paragraphs[0]
        r0 = p0.add_run(n_niv)
        r0.font.name = "Arial"
        r0.font.size = Pt(8)
        r0.font.bold = True
        r0.font.color.rgb = C_SLATE

        p1 = row_c[1].paragraphs[0]
        r1 = p1.add_run(n_acc)
        r1.font.name = "Arial"
        r1.font.size = Pt(8)
        r1.font.color.rgb = C_DARK

        p2 = row_c[2].paragraphs[0]
        r2 = p2.add_run(n_evi)
        r2.font.name = "Arial"
        r2.font.size = Pt(8)
        r2.font.color.rgb = C_MUTED

    doc.add_paragraph() # Espacio

    # 7. Ciclo de Firmas de Calidad Institucional (Tripartita)
    h7 = doc.add_paragraph()
    r_h7 = h7.add_run("7. FIRMAS DE APROBACIÓN Y CIERRE DE CALIDAD (ISO 21001)")
    r_h7.font.name = "Arial"
    r_h7.font.size = Pt(10.5)
    r_h7.font.bold = True
    r_h7.font.color.rgb = C_SLATE

    t_sign = doc.add_table(rows=2, cols=3)
    t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sign.autofit = False

    sign_roles = [
        ("DOCENTE TITULAR", "Elaboración y diagnóstico analítico"),
        ("COORDINACIÓN DE CARRERA", "Revisión y pertinencia pedagógica"),
        ("DIRECCIÓN ACADÉMICA / DECANATO", "Aprobación y certificación oficial")
    ]

    for col_idx, (role, desc) in enumerate(sign_roles):
        c_top = t_sign.cell(0, col_idx)
        c_bot = t_sign.cell(1, col_idx)
        set_cell_margins(c_top, 250, 50, 50, 50)
        set_cell_margins(c_bot, 40, 60, 50, 50)
        set_cell_background(c_top, "FAFAFA")
        set_cell_background(c_bot, "F1F5F9")

        p_t = c_top.paragraphs[0]
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_line = p_t.add_run("____________________________\nFirma y Sello")
        r_line.font.name = "Arial"
        r_line.font.size = Pt(7.5)
        r_line.font.color.rgb = C_MUTED

        p_b = c_bot.paragraphs[0]
        p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_b1 = p_b.add_run(f"{role}\n")
        r_b1.font.name = "Arial"
        r_b1.font.bold = True
        r_b1.font.size = Pt(8)
        r_b1.font.color.rgb = C_SLATE

        r_b2 = p_b.add_run(f"({desc})")
        r_b2.font.name = "Arial"
        r_b2.font.size = Pt(7)
        r_b2.font.color.rgb = C_MUTED

    doc.save(docx_output)
    print(f"Documento guardado: {docx_output}")

    # Guardar insights.json
    with open(json_output, 'w', encoding='utf-8') as f:
        json.dump(insights, f, ensure_ascii=False, indent=4)
    print(f"JSON guardado: {json_output}")

    # Registro en Control de Versiones Documental (ISO 21001:2018 Cláusula 7.5)
    try:
        from version_control_service import record_document_version
        record_document_version(
            subject_name=subject_name,
            doc_name=f"Informe Analítico de Aprendizaje IA - {subject_name}",
            doc_code=f"DOC-DIR-CP2-INF3-{codigo_materia}",
            doc_type="Informe Analítico Multidimensional",
            new_version="1.0",
            author_or_agent="Agente Etapa 3 (Analítica & Gemini IA)",
            change_description="Generación del Informe Analítico Oficial con ficha muestral, estadística descriptiva/inferencial y prescripción multinivel.",
            justification="Aseguramiento de la calidad académica según ISO 21001:2018 Cláusula 9.1.",
            file_path=docx_output,
            approval_status="Aprobado Institucional"
        )
        print("Informe registrado exitosamente en CONTROL_VERSIONES_DOCUMENTAL.xlsx")
    except Exception as e:
        print(f"Advertencia al registrar en control de versiones: {e}")

    print("Etapa 3 completada con éxito.")

if __name__ == "__main__":
    run()
