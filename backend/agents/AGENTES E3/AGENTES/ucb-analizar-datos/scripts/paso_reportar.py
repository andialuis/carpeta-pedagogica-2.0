"""Paso: reportar

Escribe lo que esta etapa entrega:

  - RESULTADOS_ANALISIS.xlsx con sus ocho hojas, que es lo que va a consumir la
    etapa de visualizacion sin tener que recalcular ni una cifra;
  - los graficos que respaldan cada hallazgo, guardados como imagen;
  - INFORME_FINAL_ANALISIS.docx con sus cuatro apartados;
  - una huella del contenido, para comparar resultados entre computadoras.

El Word es para leer y el Excel es para que la etapa siguiente lo use. Los dos
salen de las mismas listas, en la misma corrida: no pueden decir cosas distintas.

Los graficos de aqui son imagenes fijas de respaldo. El tablero interactivo es de
la etapa siguiente.
"""

import datetime as _dt
import hashlib
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ENCABEZADO = PatternFill("solid", fgColor="003366")
LETRA_ENCABEZADO = Font(color="FFFFFF", bold=True)
FILA_PAR = PatternFill("solid", fgColor="F2F4F8")


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, _dt.datetime):
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(valor, _dt.date):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def _celda(valor):
    """openpyxl no sabe escribir listas ni diccionarios: se pasan como texto."""
    if isinstance(valor, (list, tuple, dict, set)):
        return _texto(valor)
    if isinstance(valor, (int, float, str, type(None), _dt.datetime, _dt.date)):
        return valor
    return _texto(valor)


def _escribir_hoja(libro, titulo, campos, registros):
    hoja = libro.create_sheet(titulo[:31])
    hoja.append(list(campos))
    for celda in hoja[1]:
        celda.fill = ENCABEZADO
        celda.font = LETRA_ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for numero, registro in enumerate(registros, start=2):
        hoja.append([_celda(registro.get(c, "")) for c in campos])
        if numero % 2 == 0:
            for celda in hoja[numero]:
                celda.fill = FILA_PAR
    if not registros:
        hoja.append(["(sin registros)"] + [""] * (len(campos) - 1))
    hoja.freeze_panes = "A2"
    for indice, campo in enumerate(campos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = \
            max(14, min(52, len(str(campo)) + 6))
    return hoja


HOJAS = [
    ("HALLAZGOS", ["hallazgo", "cifra", "estudiantes_afectados",
                   "si_no_hago_nada", "en_que_se_apoya", "prioridad",
                   "grafico"]),
    ("HIPOTESIS", ["hipotesis", "origen", "veredicto", "cifra", "estudiantes",
                   "p", "efecto", "sobrevive_al_ajuste", "detalle"]),
    ("PERFILES", ["perfil", "estudiantes", "codigos", "rendimiento_medio",
                  "esfuerzo_medio", "abandonos_en_el_perfil", "que_significa"]),
    ("RELACIONES", ["columna_a", "columna_b", "relacion", "fuerza", "sentido",
                    "estudiantes", "p", "sobrevive_al_ajuste", "tipo",
                    "lectura"]),
    ("INDICADORES", ["que_medir", "se_enciende_cuando", "por_que_ese_valor",
                     "que_hacer", "estudiantes_hoy"]),
    ("ACCIONES", ["prioridad", "para_quien", "cuantos", "codigos", "que_hacer",
                  "en_que_hallazgo_se_apoya", "como_sabre_si_funciono"]),
    ("PRIORIDADES", ["puesto", "codigo", "urgencia", "situacion", "perfil",
                     "senales_disponibles", "motivo"]),
    ("DATOS_GRAFICOS", ["grafico", "etiqueta", "valor", "serie"]),
]


def escribir_resultados(ruta, resultados):
    """RESULTADOS_ANALISIS.xlsx: ocho hojas, mas las de contexto del analisis."""
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, campos in HOJAS:
        _escribir_hoja(libro, nombre, campos, resultados.get(nombre, []))

    # Estas dos no las pidio el formato, pero sin ellas el docente no puede
    # auditar por que el programa analizo unas columnas y otras no.
    _escribir_hoja(libro, "COLUMNAS", ["columna", "papel", "por_que", "con_dato",
                                       "distintos"],
                   resultados.get("COLUMNAS", []))
    _escribir_hoja(libro, "NO_SE_PUDO", ["analisis", "por_que"],
                   resultados.get("NO_SE_PUDO", []))
    libro.save(ruta)
    return ruta


def huella(resultados):
    """Huella del contenido analitico, para comparar corridas entre computadoras.

    No incluye las imagenes: dos maquinas con la misma version de matplotlib
    pueden generar archivos distintos byte a byte y decir exactamente lo mismo.
    """
    piezas = []
    for nombre, campos in HOJAS:
        for registro in resultados.get(nombre, []):
            piezas.append("|".join(_texto(registro.get(c, "")) for c in campos))
    return hashlib.sha256("\n".join(piezas).encode("utf-8")).hexdigest()[:16]


# --------------------------------------------------------------------- graficos
def dibujar(carpeta, resultados, config):
    """Genera las imagenes de respaldo. Devuelve [(clave, ruta, titulo, pie)]."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return []

    colores = config["colores"]
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    imagenes = []

    def guardar(figura, nombre):
        ruta = carpeta / nombre
        figura.tight_layout()
        figura.savefig(ruta, dpi=150, metadata={"Software": ""})
        plt.close(figura)
        return ruta

    def estilo(ejes, titulo):
        ejes.set_title(titulo, color=colores["ucb"], fontsize=11, loc="left")
        ejes.spines["top"].set_visible(False)
        ejes.spines["right"].set_visible(False)
        ejes.grid(axis="y", color="#DDE1E7", linewidth=0.7)
        ejes.set_axisbelow(True)

    # 1. Como se reparte el rendimiento del curso.
    valores = [s["rendimiento"] for s in resultados["_senales"]
               if s["rendimiento"] is not None]
    if valores:
        figura, ejes = plt.subplots(figsize=(6.2, 3.1))
        ejes.hist(valores, bins=min(10, max(4, len(valores) // 3)),
                  color=colores["ucb"], edgecolor="white")
        estilo(ejes, "Como se reparte el rendimiento del curso")
        ejes.set_xlabel("rendimiento (0 a 1)")
        ejes.set_ylabel("estudiantes")
        imagenes.append(("rendimiento", guardar(figura, "01_rendimiento.png"),
                         "Distribucion del rendimiento",
                         f"{len(valores)} estudiantes con datos suficientes"))

    # 2. Cuantos estudiantes hay en cada perfil.
    perfiles = resultados.get("PERFILES", [])
    if perfiles:
        figura, ejes = plt.subplots(figsize=(6.2, 3.4))
        nombres = [p["perfil"] for p in perfiles][::-1]
        cantidades = [p["estudiantes"] for p in perfiles][::-1]
        barras = ejes.barh(nombres, cantidades, color=colores["apoyo"])
        if barras:
            barras[-1].set_color(colores["alerta"])
        estilo(ejes, "Cuantos estudiantes hay en cada perfil")
        ejes.set_xlabel("estudiantes")
        ejes.grid(axis="x", color="#DDE1E7", linewidth=0.7)
        imagenes.append(("perfiles", guardar(figura, "02_perfiles.png"),
                         "Perfiles ocultos del curso",
                         "cada barra es un grupo con un comportamiento distinto"))

    # 3. La relacion mas fuerte que se sostuvo.
    dispersion = resultados.get("_dispersion")
    if dispersion:
        figura, ejes = plt.subplots(figsize=(6.2, 3.4))
        ejes.scatter(dispersion["x"], dispersion["y"], s=42,
                     color=colores["ucb"], alpha=0.8, edgecolor="white")
        estilo(ejes, dispersion["titulo"])
        ejes.set_xlabel(dispersion["etiqueta_x"][:60])
        ejes.set_ylabel(dispersion["etiqueta_y"][:60])
        imagenes.append(("relacion", guardar(figura, "03_relacion.png"),
                         dispersion["titulo"],
                         f"{len(dispersion['x'])} estudiantes; cada punto es uno"))

    # 4. Vigilado contra no vigilado.
    comparacion = resultados.get("_instrumentos", {}).get("por_estudiante", [])
    if comparacion:
        figura, ejes = plt.subplots(figsize=(6.2, 3.2))
        ejes.scatter([c["vigilado"] for c in comparacion],
                     [c["no_vigilado"] for c in comparacion], s=42,
                     color=colores["acento"], edgecolor="white")
        limite = [0, 1]
        ejes.plot(limite, limite, color=colores["gris"], linewidth=1,
                  linestyle="--")
        estilo(ejes, "Evaluaciones vigiladas frente a las hechas por cuenta propia")
        ejes.set_xlabel("promedio vigilado (0 a 1)")
        ejes.set_ylabel("promedio no vigilado (0 a 1)")
        imagenes.append(("instrumentos", guardar(figura, "04_instrumentos.png"),
                         "Vigilado frente a no vigilado",
                         "los puntos muy por encima de la linea rindieron mucho "
                         "mejor sin vigilancia"))

    # 5. A que hora trabaja el curso.
    actividad = resultados.get("_actividad", {})
    if actividad.get("disponible") and actividad.get("por_hora"):
        figura, ejes = plt.subplots(figsize=(6.6, 3.0))
        horas = [h["hora"] for h in actividad["por_hora"]]
        eventos = [h["eventos"] for h in actividad["por_hora"]]
        colores_barra = [colores["alerta"] if i < 6 else colores["ucb"]
                         for i in range(24)]
        ejes.bar(horas, eventos, color=colores_barra)
        estilo(ejes, "A que hora trabaja el curso en el aula virtual")
        ejes.set_ylabel("eventos")
        ejes.tick_params(axis="x", labelrotation=90, labelsize=7)
        imagenes.append(("actividad", guardar(figura, "05_actividad.png"),
                         "Linea de tiempo del aula virtual",
                         "en rojo, la actividad de madrugada"))

    return imagenes


def datos_de_graficos(resultados, imagenes):
    """La hoja DATOS_GRAFICOS: los numeros exactos que hay detras de cada imagen."""
    filas = []
    for s in resultados["_senales"]:
        if s["rendimiento"] is not None:
            filas.append({"grafico": "rendimiento", "etiqueta": s["codigo"],
                          "valor": s["rendimiento"], "serie": "rendimiento"})
    for p in resultados.get("PERFILES", []):
        filas.append({"grafico": "perfiles", "etiqueta": p["perfil"],
                      "valor": p["estudiantes"], "serie": "estudiantes"})
    dispersion = resultados.get("_dispersion")
    if dispersion:
        for codigo, x, y in zip(dispersion["codigos"], dispersion["x"],
                                dispersion["y"]):
            filas.append({"grafico": "relacion", "etiqueta": codigo, "valor": x,
                          "serie": dispersion["etiqueta_x"]})
            filas.append({"grafico": "relacion", "etiqueta": codigo, "valor": y,
                          "serie": dispersion["etiqueta_y"]})
    for c in resultados.get("_instrumentos", {}).get("por_estudiante", []):
        filas.append({"grafico": "instrumentos", "etiqueta": c["codigo"],
                      "valor": c["vigilado"], "serie": "vigilado"})
        filas.append({"grafico": "instrumentos", "etiqueta": c["codigo"],
                      "valor": c["no_vigilado"], "serie": "no vigilado"})
    actividad = resultados.get("_actividad", {})
    for h in actividad.get("por_hora", []):
        filas.append({"grafico": "actividad", "etiqueta": h["hora"],
                      "valor": h["eventos"], "serie": "eventos por hora"})
    return filas


# ----------------------------------------------------------------- informe Word
def escribir_informe(ruta, materia, contenido, imagenes):
    """INFORME_FINAL_ANALISIS.docx con sus cuatro apartados."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    UCB = RGBColor(0x00, 0x33, 0x66)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # --- portada -----------------------------------------------------------
    titulo = doc.add_heading("Informe de analisis del aprendizaje", level=0)
    for corrida in titulo.runs:
        corrida.font.color.rgb = UCB
    subtitulo = doc.add_paragraph()
    corrida = subtitulo.add_run(f"{materia}   ·   Etapa 3 — Analizar")
    corrida.font.size = Pt(14)
    corrida.font.color.rgb = UCB
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(
        f"Generado el {_dt.date.today().strftime('%d/%m/%Y')}. "
        "Todas las cifras se calcularon sobre el curso completo, activos y "
        "abandonos juntos. Los estudiantes aparecen por su codigo: esta etapa no "
        "tiene acceso a los nombres.")

    def tabla(campos, registros, limite=60):
        if not registros:
            doc.add_paragraph("Sin registros en este apartado.")
            return
        t = doc.add_table(rows=1, cols=len(campos))
        t.style = "Light Grid Accent 1"
        for celda, nombre in zip(t.rows[0].cells, campos):
            celda.text = str(nombre)
            for parrafo in celda.paragraphs:
                for corrida in parrafo.runs:
                    corrida.bold = True
        for registro in registros[:limite]:
            fila = t.add_row().cells
            for celda, campo in zip(fila, campos):
                celda.text = _texto(registro.get(campo, ""))[:220]
        if len(registros) > limite:
            doc.add_paragraph(
                f"Se muestran {limite} de {len(registros)} filas. La lista "
                "completa esta en RESULTADOS_ANALISIS.xlsx.")

    def bloque_hallazgo(numero, hallazgo, imagen):
        encabezado = doc.add_heading(
            f"{numero}. {hallazgo['hallazgo']}", level=2)
        for corrida in encabezado.runs:
            corrida.font.color.rgb = UCB
        for etiqueta, campo in (("La cifra que lo respalda", "cifra"),
                                ("A cuantos estudiantes afecta",
                                 "estudiantes_afectados"),
                                ("Que pasa si no hago nada", "si_no_hago_nada")):
            parrafo = doc.add_paragraph()
            corrida = parrafo.add_run(f"{etiqueta}: ")
            corrida.bold = True
            parrafo.add_run(_texto(hallazgo.get(campo, "")))
        apoyo = doc.add_paragraph()
        corrida = apoyo.add_run("En que se apoya: ")
        corrida.italic = True
        detalle = apoyo.add_run(_texto(hallazgo.get("en_que_se_apoya", "")))
        detalle.italic = True
        detalle.font.size = Pt(9)
        if imagen:
            doc.add_picture(str(imagen[1]), width=Inches(5.6))
            pie = doc.add_paragraph(imagen[3])
            pie.runs[0].font.size = Pt(8)
            pie.runs[0].font.color.rgb = RGBColor(0x7A, 0x82, 0x90)

    # --- 1. Lo que descubrimos --------------------------------------------
    if contenido.get("sospecha"):
        parrafo = doc.add_paragraph()
        corrida = parrafo.add_run("Lo que el docente queria averiguar: ")
        corrida.bold = True
        parrafo.add_run(contenido["sospecha"])

    doc.add_heading("Lo que descubrimos", level=1)
    doc.add_paragraph(contenido["panorama"])
    if contenido.get("avisos"):
        doc.add_heading("Antes de leer: lo que hay que tener en cuenta", level=2)
        for aviso in contenido["avisos"]:
            parrafo = doc.add_paragraph(str(aviso), style="List Bullet")
            for corrida in parrafo.runs:
                corrida.font.size = Pt(10)
    por_clave = {i[0]: i for i in imagenes}
    if "rendimiento" in por_clave:
        imagen = por_clave.pop("rendimiento")
        doc.add_picture(str(imagen[1]), width=Inches(5.6))
        pie = doc.add_paragraph(imagen[3])
        pie.runs[0].font.size = Pt(8)
        pie.runs[0].font.color.rgb = RGBColor(0x7A, 0x82, 0x90)

    usadas = set()
    for numero, hallazgo in enumerate(contenido["hallazgos"], start=1):
        clave = hallazgo.get("grafico", "")
        imagen = None
        if clave and clave in por_clave and clave not in usadas:
            imagen = por_clave[clave]
            usadas.add(clave)
        bloque_hallazgo(numero, hallazgo, imagen)
    if contenido["faltantes"]:
        parrafo = doc.add_paragraph()
        corrida = parrafo.add_run(
            f"Faltaron {contenido['faltantes']} hallazgo(s) para llegar a "
            f"{contenido['objetivo']}. ")
        corrida.bold = True
        parrafo.add_run(
            "No se completaron con conclusiones que los datos no sostienen. "
            "Para conseguirlos haria falta: "
            + "; ".join(contenido["que_haria_falta"]) + ".")

    if contenido.get("sobrantes"):
        doc.add_paragraph(
            f"Otros {contenido['sobrantes']} hallazgo(s) se sostienen pero no "
            "entraron en este informe para que se pueda leer. Estan completos "
            "en la hoja HALLAZGOS de RESULTADOS_ANALISIS.xlsx.")

    doc.add_heading("Los grupos que el promedio esconde", level=2)
    tabla(["perfil", "estudiantes", "rendimiento_medio", "que_significa"],
          contenido["perfiles"])

    # --- 2. Verificacion de sospechas -------------------------------------
    doc.add_heading("Verificacion de sospechas", level=1)
    doc.add_paragraph(
        "Una sospecha descartada tambien es un resultado: evita gastar el "
        "semestre corrigiendo un problema que no existe.")
    tabla(["hipotesis", "origen", "veredicto", "cifra", "estudiantes", "detalle"],
          contenido["hipotesis"])

    doc.add_heading("Relaciones entre variables", level=2)
    if contenido["aviso_tamano"]:
        aviso = doc.add_paragraph(contenido["aviso_tamano"])
        aviso.runs[0].bold = True
    tabla(["columna_a", "columna_b", "relacion", "fuerza", "sentido",
           "estudiantes", "tipo", "lectura"], contenido["relaciones"])
    if contenido["redundantes"]:
        doc.add_heading("Cruces que se apartaron por triviales", level=3)
        doc.add_paragraph(
            "Estas parejas se mueven juntas por construccion, no por un "
            "hallazgo. Se apartan para que no ocupen el lugar de una relacion real.")
        tabla(["columna_a", "columna_b", "relacion", "por_que_se_aparta"],
              contenido["redundantes"], limite=25)

    # --- 3. Mis proximos pasos --------------------------------------------
    doc.add_heading("Mis proximos pasos", level=1)
    doc.add_paragraph(
        "Una accion por grupo. Las escribio esta etapa; la etapa de "
        "visualizacion solo las muestra y tiene prohibido inventar otras.")
    tabla(["prioridad", "para_quien", "cuantos", "que_hacer",
           "como_sabre_si_funciono"], contenido["acciones"])

    doc.add_heading("Senales de alerta para el proximo semestre", level=2)
    tabla(["que_medir", "se_enciende_cuando", "que_hacer", "estudiantes_hoy"],
          contenido["indicadores"])

    # --- 4. Conclusiones tecnicas -----------------------------------------
    doc.add_heading("Conclusiones tecnicas", level=1)
    doc.add_paragraph(contenido["conclusion_tecnica"])
    doc.add_heading("Lo que no se pudo analizar", level=2)
    tabla(["analisis", "por_que"], contenido["no_se_pudo"])
    doc.add_heading("Como leyo el programa cada columna", level=2)
    doc.add_paragraph(
        "Si alguna columna quedo con el papel equivocado, el analisis completo "
        "cambia. Esta tabla existe para poder revisarlo.")
    tabla(["columna", "papel", "por_que"], contenido["columnas"], limite=80)

    doc.save(ruta)
    return ruta
