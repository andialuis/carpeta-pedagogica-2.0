"""Paso: leer_depurada

Abre la BASE_DATOS_DEPURADA.xlsx que dejo la Etapa 2 y averigua QUE ES cada
columna, sin saber nada de la materia.

Este es el paso del que depende que la habilidad sirva para cualquier curso. El
programa no trae escrito ni un nombre de columna, ni un instrumento, ni una escala:
reconoce el papel de cada columna por como se comporta y por las palabras de
`config.json`, que el docente puede ampliar si su materia usa otro vocabulario.

Si a la base le falta una hoja, se sigue adelante sin ella y queda anotado. Lo que
nunca se hace es inventar una columna para tapar el hueco.
"""

import datetime as _dt
import unicodedata

import openpyxl

VACIO, NUMERO, TEXTO, FECHA = "vacio", "numero", "texto", "fecha"

# Papeles que puede tener una columna. El resto del programa decide que hacer con
# cada una mirando esta etiqueta, nunca su nombre.
IDENTIDAD = "identidad protegida"
NOTA = "calificacion"
ASISTENCIA = "asistencia"
PLATAFORMA = "actividad en plataforma"
DURACION = "tiempo empleado"
MOMENTO = "fecha u hora"
CONTEXTO = "contexto del estudiante"
DERIVADA = "derivada de otra columna"
CONSTANTE = "vale lo mismo para todos"
SIN_EMPAREJAR = "no se emparejo con ningun estudiante"
PAPELEO = "papeleo de la transcripcion"
OTRO = "sin clasificar"


def tipo_de(valor):
    """Que clase de dato es una celda, sin convertirla."""
    if valor is None:
        return VACIO
    if isinstance(valor, bool):
        return TEXTO
    if isinstance(valor, (_dt.datetime, _dt.date, _dt.time)):
        return FECHA
    if isinstance(valor, (int, float)):
        return NUMERO
    texto = str(valor).strip()
    if not texto:
        return VACIO
    try:
        float(texto.replace(",", "."))
        return NUMERO
    except ValueError:
        return TEXTO


def a_numero(valor):
    """Devuelve el valor como numero, o None si no lo es."""
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    if valor is None:
        return None
    try:
        return float(str(valor).strip().replace(",", "."))
    except ValueError:
        return None


def a_momento(valor):
    """Devuelve la celda como fecha y hora, o None."""
    if isinstance(valor, _dt.datetime):
        return valor
    if isinstance(valor, _dt.date):
        return _dt.datetime(valor.year, valor.month, valor.day)
    if isinstance(valor, str):
        texto = valor.strip()
        for formato in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M:%S",
                        "%d/%m/%Y %H:%M", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return _dt.datetime.strptime(texto, formato)
            except ValueError:
                continue
    return None


def normalizar(texto):
    """Minusculas, sin tildes y sin caracteres invisibles.

    Sirve para comparar nombres de columnas sin que una tilde o una mayuscula
    decidan si una materia se analiza bien o mal.
    """
    if texto is None:
        return ""
    plano = unicodedata.normalize("NFKD", str(texto))
    plano = "".join(c for c in plano if not unicodedata.combining(c))
    return " ".join(plano.lower().replace("​", " ").split())


def _contiene(nombre, palabras):
    """¿El titulo contiene alguna de esas palabras?

    Los guiones bajos y los guiones se tratan como espacios. Un archivo
    exportado del aula casi siempre se llama «aula_virtual_accesos» o
    «moodle-log», y buscar «aula virtual» con espacio no lo encontraba nunca.
    Era la razon por la que la linea de tiempo no se armaba en cursos donde el
    archivo estaba ahi mismo.
    """
    plano = normalizar(nombre).replace("_", " ").replace("-", " ")
    plano = " ".join(plano.split())
    return any(p in plano for p in palabras)


def _sin_sufijo(nombre):
    """Quita el sufijo entre corchetes que agrega la Etapa 2: '... [sobre 6]'."""
    plano = str(nombre or "").rstrip()
    if plano.endswith("]") and "[" in plano:
        corte = plano.rfind("[")
        # El prefijo [archivo] va al principio; el sufijo derivado, al final.
        if corte > 0:
            return plano[:corte].rstrip()
    return plano


def _origen(nombre):
    """El archivo del que vino la columna, si la Etapa 1 lo dejo entre corchetes."""
    plano = str(nombre or "").lstrip()
    if plano.startswith("[") and "]" in plano:
        return plano[1:plano.index("]")].strip()
    return ""


# ------------------------------------------------------------------- lectura
def _leer_hoja(hoja):
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return [], []
    columnas = [c if c is not None else f"columna_{i + 1}"
                for i, c in enumerate(filas[0])]
    cuerpo = [list(f) + [None] * (len(columnas) - len(f)) for f in filas[1:]]
    cuerpo = [f for f in cuerpo if any(tipo_de(v) != VACIO for v in f)]
    return columnas, cuerpo


def leer(ruta):
    """Devuelve las hojas de la base depurada tal como vinieron."""
    libro = openpyxl.load_workbook(ruta, data_only=True)
    datos = {"hojas_presentes": list(libro.sheetnames), "avisos": []}

    for nombre in ("ACTIVOS", "ABANDONOS"):
        if nombre in libro.sheetnames:
            columnas, filas = _leer_hoja(libro[nombre])
        else:
            columnas, filas = [], []
            datos["avisos"].append(
                f"la base no trae la hoja {nombre}; se sigue sin ella")
        datos[nombre.lower()] = {"columnas": columnas, "filas": filas}

    reglas = []
    if "REGLAS_EVALUACION" in libro.sheetnames:
        columnas, filas = _leer_hoja(libro["REGLAS_EVALUACION"])
        reglas = [dict(zip(columnas, f)) for f in filas]
    else:
        datos["avisos"].append(
            "la base no trae REGLAS_EVALUACION: la medida de rendimiento se "
            "construira promediando los instrumentos, sin ponderar")
    datos["reglas"] = reglas

    # La hoja ANOMALIAS de la etapa anterior no se corrige aqui (R1), pero sirve
    # para avisar al docente de que su base todavia arrastra valores sospechosos.
    datos["anomalias_heredadas"] = []
    if "ANOMALIAS" in libro.sheetnames:
        columnas, filas = _leer_hoja(libro["ANOMALIAS"])
        datos["anomalias_heredadas"] = [dict(zip(columnas, f)) for f in filas]

    libro.close()
    return datos


def unir_curso(datos):
    """Junta ACTIVOS y ABANDONOS en una sola tabla, con su procedencia.

    Los que abandonaron no son datos descartados: suelen ser la mitad de la
    respuesta. Se analizan juntos y cada fila recuerda de que hoja vino.
    """
    columnas = list(datos["activos"]["columnas"])
    filas = [list(f) for f in datos["activos"]["filas"]]
    situacion = ["activo"] * len(filas)

    columnas_ab = datos["abandonos"]["columnas"]
    if columnas_ab and datos["abandonos"]["filas"]:
        # La hoja ABANDONOS de la Etapa 2 trae solo el codigo y el motivo. Se
        # completan sus columnas con vacio para que la tabla siga siendo una sola.
        posicion = {normalizar(c): i for i, c in enumerate(columnas)}
        for fila_ab in datos["abandonos"]["filas"]:
            nueva = [None] * len(columnas)
            for nombre, valor in zip(columnas_ab, fila_ab):
                indice = posicion.get(normalizar(nombre))
                if indice is not None:
                    nueva[indice] = valor
            filas.append(nueva)
            situacion.append("abandono")

    columnas.append("situacion_en_el_curso")
    for fila, estado in zip(filas, situacion):
        fila.append(estado)

    # La hoja ABANDONOS de la Etapa 2 guarda el motivo de la separacion, no las
    # calificaciones de esas personas. Si es asi, entran en el conteo del curso
    # pero no en los cruces, y el docente tiene que saberlo: creer que se
    # analizo a los que se fueron cuando no habia con que es peor que no
    # analizarlos.
    aviso = ""
    if datos["abandonos"]["filas"]:
        utiles = [c for c in columnas_ab
                  if normalizar(c) not in ("codigo", "motivo", "desde",
                                           "senales", "quien_decidio")]
        if len(utiles) <= 1:
            aviso = (
                f"los {len(datos['abandonos']['filas'])} estudiante(s) que "
                "abandonaron entran en el conteo del curso, pero la hoja "
                "ABANDONOS solo trae el motivo de su separacion y no sus "
                "calificaciones ni su asistencia: no se los puede incluir en "
                "los cruces ni en los perfiles. Para analizarlos hace falta que "
                "la Etapa 2 escriba esa hoja con todas las columnas.")
    return columnas, filas, aviso


def columna_clave(columnas, filas):
    """La columna que identifica al estudiante: la del codigo de la Etapa 2."""
    for indice, nombre in enumerate(columnas):
        if normalizar(nombre) in ("codigo", "codigo_estudiante", "id"):
            return indice
    # Si no se llama 'codigo', la clave es la primera columna cuyos valores no se
    # repiten nunca.
    for indice in range(len(columnas)):
        valores = [str(f[indice]).strip() for f in filas
                   if tipo_de(f[indice]) != VACIO]
        if valores and len(set(valores)) == len(filas):
            return indice
    return 0


# -------------------------------------------------------- papel de cada columna
def clasificar_columnas(columnas, filas, clave, config, reglas=()):
    """Decide el papel de cada columna. Devuelve una lista de diccionarios.

    El orden de las reglas importa: primero lo que se reconoce por como se
    comporta (identidad protegida, constante, derivada) y solo despues lo que se
    reconoce por el vocabulario del titulo.
    """
    palabras = config["palabras"]
    maximo_categorias = config["maximo_categorias_contexto"]
    minimo = config["minimo_datos_para_estadistica"]

    codigos = [str(f[clave]).strip() for f in filas]
    nombres_normalizados = {normalizar(c): i for i, c in enumerate(columnas)}
    instrumentos = {normalizar(r.get("instrumento", "")) for r in reglas
                    if r.get("instrumento")}

    fichas = []
    for indice, nombre in enumerate(columnas):
        valores = [f[indice] for f in filas]
        con_dato = [v for v in valores if tipo_de(v) != VACIO]
        numeros = [a_numero(v) for v in con_dato]
        numeros = [n for n in numeros if n is not None]
        distintos = {str(v).strip() for v in con_dato}

        ficha = {
            "columna": nombre,
            "indice": indice,
            "origen": _origen(nombre),
            "con_dato": len(con_dato),
            "distintos": len(distintos),
            "es_numerica": bool(con_dato) and len(numeros) >= len(con_dato) * 0.8,
            "papel": OTRO,
            "por_que": "",
        }

        if indice == clave:
            ficha.update(papel=IDENTIDAD,
                         por_que="es el codigo con el que se identifica a cada "
                                 "estudiante")
            fichas.append(ficha)
            continue

        # 1. Identidad que la Etapa 2 ya protegio: la columna repite el codigo.
        #    No es una columna sin variacion, es una columna anonimizada, y
        #    confundirlas seria acusar a la etapa anterior de haber roto los datos.
        if con_dato and all(str(v).strip() == c
                            for v, c in zip(valores, codigos)
                            if tipo_de(v) != VACIO):
            ficha.update(papel=IDENTIDAD,
                         por_que="la etapa anterior reemplazo su contenido por el "
                                 "codigo del estudiante para proteger la identidad")
            fichas.append(ficha)
            continue

        if _contiene(nombre, palabras["no_analizable"]) and not ficha["es_numerica"]:
            ficha.update(papel=IDENTIDAD,
                         por_que="su titulo indica que identifica a la persona")
            fichas.append(ficha)
            continue

        # 2. Papeleo de una transcripcion asistida: de que archivo salio la fila
        # y con cuanta seguridad se leyo. Es trazabilidad, no un dato del
        # estudiante. Sin esto, «confianza» se convertia en una variable de
        # contexto y el programa llegaba a comparar el rendimiento de los
        # estudiantes segun lo seguro que estuvo quien transcribio la planilla.
        if _contiene(nombre, ("archivo origen", "confianza", "origen")) and \
                "transcripcion" in normalizar(nombre):
            ficha.update(papel=PAPELEO,
                         por_que="dice de que archivo salio el dato y con cuanta "
                                 "seguridad se leyo, no es una caracteristica "
                                 "del estudiante")
            fichas.append(ficha)
            continue

        # 3. Una columna sin un solo dato no es una columna constante: es un
        # empalme que fallo. Llamarla «vale lo mismo para todos» dejaba al
        # docente creyendo que su transcripcion entro cuando en realidad no
        # engancho con nadie, y el aviso se perdia entre las columnas normales.
        if not con_dato:
            ficha.update(papel=SIN_EMPAREJAR,
                         por_que="no tiene ni un dato: ninguna fila de ese archivo "
                                 "encontro a su estudiante. Revise que la columna "
                                 "de identidad use el mismo codigo o el mismo "
                                 "nombre que la lista oficial")
            fichas.append(ficha)
            continue

        # 4. Sin variacion: no se puede analizar, pero se dice en vez de ignorarla.
        if len(distintos) <= 1:
            ficha.update(papel=CONSTANTE,
                         por_que="todos los estudiantes tienen el mismo valor, asi "
                                 "que no puede explicar ninguna diferencia entre ellos")
            fichas.append(ficha)
            continue

        # 3. Derivada: la Etapa 2 la creo a partir de otra que sigue en la base.
        base = _sin_sufijo(nombre)
        if base != str(nombre).rstrip() and normalizar(base) in nombres_normalizados:
            ficha.update(papel=DERIVADA, derivada_de=base,
                         por_que=f"es la version convertida de «{base}», que sigue "
                                 "en la base: se mueven juntas por construccion")
            fichas.append(ficha)
            continue

        # 4. Por vocabulario del titulo.
        if _contiene(nombre, palabras["duracion"]):
            ficha.update(papel=DURACION,
                         por_que="su titulo habla de cuanto tiempo tardo el estudiante")
        elif _contiene(nombre, palabras["hora_inicio"]) or (
                con_dato and all(a_momento(v) is not None for v in con_dato[:10])):
            ficha.update(papel=MOMENTO,
                         por_que="contiene fechas u horas, no una medida del estudiante")
        elif _contiene(nombre, palabras["asistencia"]):
            # La asistencia se deja tal cual aunque venga en texto: la Etapa 2
            # crea al lado su gemela numerica, y la segunda pasada de mas abajo
            # se queda con la que se puede sumar.
            ficha.update(papel=ASISTENCIA,
                         por_que="su titulo corresponde a un registro de asistencia")
        elif _contiene(nombre, palabras["plataforma"]) and not ficha["es_numerica"] \
                and 2 <= len(distintos) <= maximo_categorias:
            # El titulo habla del aula virtual, pero el contenido son categorias
            # y no numeros: «franja horaria mas frecuente» dice «madrugada» o
            # «tarde». Antes ganaba el titulo y la columna quedaba entre dos
            # sillas: como medida no se puede promediar, y como grupo no se la
            # ofrecia para comparar. La materia perdia a la vez la senal y la
            # unica variable que permitia partir el curso en dos. A diferencia de
            # la asistencia, aqui no hay ninguna gemela numerica que la rescate.
            ficha.update(papel=CONTEXTO,
                         por_que=f"habla del aula virtual pero agrupa a los "
                                 f"estudiantes en {len(distintos)} categorias, "
                                 "asi que sirve para comparar grupos")
        elif _contiene(nombre, palabras["plataforma"]):
            ficha.update(papel=PLATAFORMA,
                         por_que="su titulo corresponde a actividad en el aula virtual")
        elif ficha["es_numerica"] and (
                normalizar(base) in instrumentos
                or any(i and i in normalizar(nombre) for i in instrumentos)):
            ficha.update(papel=NOTA,
                         por_que="la hoja REGLAS_EVALUACION la declara como "
                                 "instrumento de evaluacion")
        elif ficha["es_numerica"] and len(numeros) >= minimo:
            ficha.update(papel=NOTA,
                         por_que="es numerica y se comporta como una calificacion")
        elif (not ficha["es_numerica"]
              and 2 <= len(distintos) <= maximo_categorias):
            ficha.update(papel=CONTEXTO,
                         por_que=f"agrupa a los estudiantes en {len(distintos)} "
                                 "categorias, asi que sirve para comparar grupos")
        else:
            ficha.update(papel=OTRO,
                         por_que="no se reconocio ni como medida ni como grupo")
        fichas.append(ficha)

    # Segunda pasada. La Etapa 2 deja el original al lado de su version
    # convertida: «Asistencia 12/04» con P/A/T, y «Asistencia 12/04 [1/0]» con
    # numeros. Apartar siempre la convertida dejaria la asistencia sin medir,
    # porque la que queda no se puede sumar. Se conserva la que sirve.
    por_nombre = {normalizar(f["columna"]): f for f in fichas}
    for ficha in fichas:
        if ficha["papel"] != DERIVADA:
            continue
        original = por_nombre.get(normalizar(ficha.get("derivada_de", "")))
        if original is None or original["es_numerica"] or not ficha["es_numerica"]:
            continue
        ficha["papel"], ficha["por_que"] = original["papel"], (
            f"es la version numerica de «{original['columna']}», que venia como "
            "texto: es la unica de las dos con la que se puede calcular")
        original["papel"] = DERIVADA
        original["derivada_de"] = ficha["columna"]
        original["por_que"] = (
            f"quedo reemplazada por «{ficha['columna']}», su version numerica: "
            "se conserva para poder auditar, no para analizar")

    return fichas


def limpiar_titulo(columna):
    """Convierte «[Laboratorios QUIM101 · Laboratorio 2] Nota s/10» en algo que
    un docente pueda leer en una frase: «Laboratorio 2 · Nota s/10».

    El prefijo entre corchetes dice de que archivo vino la columna. Sirve para
    auditar, no para escribirlo dentro de una hipotesis.
    """
    texto = str(columna or "").strip()
    if texto.startswith("[") and "]" in texto:
        dentro = texto[1:texto.index("]")]
        resto = texto[texto.index("]") + 1:].strip()
        # Si el archivo traia varias hojas, el nombre util es el de la hoja. Si
        # traia una sola, el nombre util es el del archivo: dejar solo «Nota
        # s/30» no le dice a nadie de que evaluacion se esta hablando.
        parte = dentro.split("·")[-1].strip() if "·" in dentro else dentro.strip()
        texto = f"{parte} · {resto}".strip(" ·") if parte else resto
    return " ".join(texto.split()) or str(columna or "")


def de_papel(fichas, *papeles):
    """Los indices de las columnas que tienen alguno de esos papeles."""
    return [f["indice"] for f in fichas if f["papel"] in papeles]


def serie(filas, indice):
    """Los valores numericos de una columna, con None donde no hay dato."""
    return [a_numero(f[indice]) for f in filas]
