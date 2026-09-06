"""Instalador y autodiagnostico de la Etapa 3 — Analizar.

Se llama instalar_3.py, con el numero de su etapa: cada etapa del ciclo trae el
suyo y asi nunca se confunden ni se pisan entre ellas.

Se ejecuta UNA vez, despues de copiar la carpeta al espacio de trabajo:

    python3 AGENTES/ucb-analizar-datos/instalar_3.py

Comprueba el entorno, verifica que la habilidad esta completa, la prueba de punta
a punta con un curso inventado y deja anunciada la habilidad en el AGENTS.md de
la raiz para que el agente la descubra por su cuenta.

No toca ningun archivo del docente.
"""

import subprocess
import sys
import tempfile
from pathlib import Path

AQUI = Path(__file__).resolve().parent          # .../AGENTES/ucb-analizar-datos
CARPETA_AGENTES = AQUI.parent                   # .../AGENTES
RAIZ_ESPACIO = CARPETA_AGENTES.parent           # el espacio de trabajo

VERSION_MINIMA = (3, 8)
OK, FALLO, AVISO = "  [OK]   ", "  [FALLO]", "  [AVISO]"

ANUNCIO = """
## Etapa 3 · Analizar

Cuando el docente pida **analizar, cruzar variables, buscar hallazgos, poner a
prueba una sospecha o encontrar grupos de riesgo** en su base ya depurada, o
mencione la Etapa 3 / ANALIZAR, lee `AGENTES/ucb-analizar-datos/SKILL.md`
completo y siguelo al pie de la letra, respetando sus dos pausas. Los archivos de
`scripts/` se ejecutan; no se leen ni se reescriben.
"""


def titulo(texto):
    print(f"\n{texto}\n" + "-" * len(texto))


# --------------------------------------------------------------------- entorno
def comprobar_python():
    actual = ".".join(str(n) for n in sys.version_info[:3])
    if sys.version_info < VERSION_MINIMA:
        minima = ".".join(str(n) for n in VERSION_MINIMA)
        print(f"{FALLO} Python {actual}: se necesita {minima} o superior.")
        print("          Instalalo desde https://www.python.org/downloads/")
        return False
    print(f"{OK} Python {actual}")
    return True


def comprobar_dependencias():
    faltan, todo_bien = [], True
    try:
        import openpyxl
        print(f"{OK} openpyxl {openpyxl.__version__}")
    except ImportError:
        print(f"{FALLO} falta openpyxl (para leer y escribir Excel)")
        faltan.append("openpyxl")
        todo_bien = False
    try:
        import docx  # noqa: F401
        print(f"{OK} python-docx instalado")
    except ImportError:
        print(f"{FALLO} falta python-docx (para el informe en Word)")
        faltan.append("python-docx")
        todo_bien = False
    try:
        import matplotlib
        print(f"{OK} matplotlib {matplotlib.__version__}")
    except ImportError:
        print(f"{FALLO} falta matplotlib (para los graficos del informe)")
        faltan.append("matplotlib")
        todo_bien = False
    if faltan:
        print(f"          Instalalas con:  "
              f"{Path(sys.executable).name} -m pip install {' '.join(faltan)}")
    return todo_bien


def comprobar_archivos():
    esperados = [
        AQUI / "SKILL.md",
        AQUI / "scripts" / "analizar.py",
        AQUI / "scripts" / "config.json",
        AQUI / "scripts" / "estadistica.py",
        AQUI / "referencias" / "formato-salida.md",
        AQUI / "ejemplos" / "crear_base_de_prueba.py",
    ]
    faltan = [p for p in esperados if not p.exists()]
    if faltan:
        print(f"{FALLO} la copia esta incompleta. Faltan:")
        for p in faltan:
            print(f"          {p.relative_to(AQUI)}")
        return False
    pasos = sorted(AQUI.glob("scripts/paso_*.py"))
    print(f"{OK} habilidad completa ({len(pasos)} pasos + punto de entrada)")
    return True


# ------------------------------------------------------------- autodiagnostico
HIPOTESIS_DE_PRUEBA = ("los estudiantes de modalidad virtual obtienen una nota "
                       "promedio mas baja que los presenciales")

# Marcador que no existe en ningun idioma ni en ningun comentario del codigo. Se
# entrega como si fuera la sospecha del docente y despues se busca: si aparece en
# algun script, el programa guardo un dato que no debia guardar (R6).
MARCADOR_R6 = "zqx7marcador"


def prueba_de_extremo_a_extremo():
    """Corre la habilidad entera sobre un curso inventado."""
    with tempfile.TemporaryDirectory(prefix="ucb_e3_") as temporal:
        banco = Path(temporal)
        generador = AQUI / "ejemplos" / "crear_base_de_prueba.py"
        hecho = subprocess.run([sys.executable, str(generador), str(banco)],
                               capture_output=True, text=True)
        if hecho.returncode != 0:
            print(f"{FALLO} no se pudo crear el curso de prueba:")
            print(hecho.stderr.strip()[:400])
            return False

        analizar = AQUI / "scripts" / "analizar.py"
        bases = banco / "CURSO-PRUEBA-REV" / "bases_de_datos"

        # --- 1. Modo plan: no debe escribir nada ----------------------------
        hecho = subprocess.run(
            [sys.executable, str(analizar), "CURSO-PRUEBA", "--raiz", str(banco),
             "--hipotesis", HIPOTESIS_DE_PRUEBA, "--plan"],
            capture_output=True, text=True)
        if hecho.returncode != 0:
            print(f"{FALLO} el modo plan no llego al final:")
            print(hecho.stderr.strip()[:400])
            return False
        if (bases / "RESULTADOS_ANALISIS.xlsx").exists():
            print(f"{FALLO} el modo plan escribio archivos y no debia")
            return False
        if "A · LA HIPOTESIS" not in hecho.stdout:
            print(f"{FALLO} el modo plan no mostro las cinco tablas de la pausa")
            return False
        print(f"{OK} el modo plan muestra el plan y no escribe nada")

        # --- 2. Ejecucion real ----------------------------------------------
        hecho = subprocess.run(
            [sys.executable, str(analizar), "CURSO-PRUEBA", "--raiz", str(banco),
             "--hipotesis", HIPOTESIS_DE_PRUEBA, "--sospecha", MARCADOR_R6],
            capture_output=True, text=True)
        if hecho.returncode != 0:
            print(f"{FALLO} la habilidad no llego al final:")
            print(hecho.stderr.strip()[:700])
            return False
        salida = hecho.stdout

        excel = bases / "RESULTADOS_ANALISIS.xlsx"
        if not excel.exists():
            print(f"{FALLO} no se genero RESULTADOS_ANALISIS.xlsx")
            return False

        import openpyxl
        libro = openpyxl.load_workbook(excel)
        esperadas = ["HALLAZGOS", "HIPOTESIS", "PERFILES", "RELACIONES",
                     "INDICADORES", "ACCIONES", "PRIORIDADES", "DATOS_GRAFICOS"]
        faltan = [h for h in esperadas if h not in libro.sheetnames]
        if faltan:
            print(f"{FALLO} a los resultados les faltan hojas: {faltan}")
            libro.close()
            return False
        hallazgos = libro["HALLAZGOS"].max_row - 1
        prioridades = libro["PRIORIDADES"].max_row - 1
        libro.close()
        print(f"{OK} RESULTADOS_ANALISIS.xlsx con sus 8 hojas "
              f"({hallazgos} hallazgos)")

        carpeta = banco / "CURSO-PRUEBA-REV" / "RESULTADOS" / "ETAPA 3 - ANALIZAR"
        if (carpeta / "INFORME_FINAL_ANALISIS.docx").exists():
            print(f"{OK} informe en Word generado")
        else:
            print(f"{FALLO} no se genero el informe en Word")
            return False

        graficos = sorted((carpeta / "graficos").glob("*.png"))
        if graficos:
            print(f"{OK} {len(graficos)} grafico(s) de respaldo generados")
        else:
            print(f"{AVISO} no se genero ningun grafico")

        # --- 3. Las garantias que no pueden fallar --------------------------
        total = _valor(salida, "Estudiantes analizados").strip()
        activos = _valor(salida, "· activos").strip()
        abandonos = _valor(salida, "· abandonos").strip()
        if total and activos and abandonos and \
                int(activos) + int(abandonos) == int(total):
            print(f"{OK} el curso completo entra en el analisis: "
                  f"{activos} activos + {abandonos} abandonos = {total}")
        else:
            print(f"{FALLO} el analisis no cubre a todo el curso")
            return False

        if prioridades and int(total) == prioridades:
            print(f"{OK} los {prioridades} estudiantes aparecen priorizados")
        else:
            print(f"{AVISO} la lista de prioridades no cubre a todo el curso")

        if "SI" in _valor(salida, "Base depurada sin modificar"):
            print(f"{OK} la base de la Etapa 2 quedo intacta")
        else:
            print(f"{AVISO} no se pudo confirmar que la entrada quedo intacta")

        # --- 4. R6: el programa no lleva datos adentro ----------------------
        rastro = _buscar_en_el_codigo(MARCADOR_R6)
        if rastro:
            print(f"{FALLO} quedo un dato del docente dentro del codigo: {rastro}")
            return False
        print(f"{OK} nada de lo que se entrego quedo escrito en el programa")

        huella = _valor(salida, "Huella del analisis").strip()
        if huella:
            print(f"{OK} huella del analisis generada ({huella})")
        return True


def _valor(salida, etiqueta):
    """Lee un valor del resumen de control impreso por el programa."""
    for linea in salida.splitlines():
        if linea.strip().startswith(etiqueta):
            return linea.split(etiqueta, 1)[1]
    return ""


def _buscar_en_el_codigo(palabra):
    """R6: ningun dato del docente puede quedar escrito en los scripts."""
    encontrados = []
    for ruta in sorted((AQUI / "scripts").glob("*.py")):
        texto = ruta.read_text(encoding="utf-8", errors="ignore").lower()
        if palabra.lower() in texto:
            encontrados.append(ruta.name)
    return encontrados


# ------------------------------------------------------------------ AGENTS.md
MARCA = "ucb-analizar-datos"


def anunciar_habilidad():
    destino = RAIZ_ESPACIO / "AGENTS.md"

    if not destino.exists():
        destino.write_text(
            "# Instrucciones para el agente\n\n"
            "Este espacio de trabajo contiene habilidades en `AGENTES/`. Cada una "
            "es una carpeta con un `SKILL.md` que describe un procedimiento, y "
            "programas deterministas que hacen el trabajo pesado.\n"
            + ANUNCIO, encoding="utf-8")
        print(f"{OK} AGENTS.md creado en {destino}")
        return

    actual = destino.read_text(encoding="utf-8", errors="ignore")
    if MARCA in actual:
        print(f"{OK} AGENTS.md ya menciona esta habilidad; no se toco")
        return

    with open(destino, "a", encoding="utf-8") as f:
        f.write("\n" + ANUNCIO)
    print(f"{OK} ya existia un AGENTS.md: se le AGREGO el anuncio de la Etapa 3.")
    print("          No se borro ni se cambio nada de lo que ya tenia.")


# -------------------------------------------------------------------------- main
def main():
    print("=" * 68)
    print(" INSTALACION — Etapa 3: Analizar")
    print("=" * 68)
    print(f"\nEspacio de trabajo detectado: {RAIZ_ESPACIO}")

    titulo("1. Entorno")
    if not (comprobar_python() and comprobar_dependencias()):
        print("\nCorrige lo anterior y vuelve a ejecutar este instalador.")
        sys.exit(1)

    titulo("2. Integridad de la copia")
    if not comprobar_archivos():
        print("\nVuelve a copiar la carpeta ucb-analizar-datos completa.")
        sys.exit(1)

    titulo("3. Prueba con un curso inventado")
    print("  (no se usa ningun dato real; la prueba se borra sola al terminar)")
    if not prueba_de_extremo_a_extremo():
        print("\nLa habilidad no paso la prueba. Revisa el mensaje de arriba.")
        sys.exit(1)

    titulo("4. Descubrimiento por el agente")
    anunciar_habilidad()

    print("\n" + "=" * 68)
    print(" LISTO")
    print("=" * 68)
    print("""
Esta etapa necesita el resultado de la Etapa 2. Comprueba que existe
NOMBRE_CARPETA-REV/bases_de_datos/BASE_DATOS_DEPURADA.xlsx antes de empezar.

Despues, escribele al agente en tus palabras:

     Quiero analizar los datos depurados de mi materia.

Te hara tres preguntas, te ayudara a dejar tu sospecha en una forma que se
pueda comprobar, y te mostrara el plan completo antes de analizar nada.
""")


if __name__ == "__main__":
    main()
