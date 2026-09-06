"""Paso: reportar

Escribe todo lo que esta etapa entrega:

  - BASE_DATOS_DEPURADA.xlsx con sus siete hojas;
  - EQUIVALENCIA_CODIGOS.xlsx, el unico archivo que vuelve a contener nombres;
  - INFORME_PROCESAMIENTO.docx, para leer y archivar;
  - la huella de la hoja ACTIVOS, que permite comparar resultados entre
    computadoras.
"""

import datetime as _dt
import hashlib

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ENCABEZADO = PatternFill("solid", fgColor="1F3864")
LETRA_ENCABEZADO = Font(color="FFFFFF", bold=True)
FILA_PAR = PatternFill("solid", fgColor="F2F4F8")


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, _dt.datetime):
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(valor, _dt.date):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def huella(columnas, filas, indice_clave):
    """Huella de los valores de ACTIVOS, ordenados por el codigo del estudiante.

    No incluye los nombres de las columnas: dos corridas pueden rotularlas
    distinto y tener exactamente los mismos datos.
    """
    lineas = []
    for fila in sorted(filas, key=lambda f: _texto(f[indice_clave])):
        lineas.append("".join(_texto(v) for v in fila))
    return hashlib.sha256("".join(lineas).encode("utf-8")).hexdigest()[:16]


def _escribir_hoja(libro, titulo, columnas, filas, plegar=()):
    hoja = libro.create_sheet(titulo[:31])
    hoja.append([_titulo_visible(c) for c in columnas])
    for celda in hoja[1]:
        celda.fill = ENCABEZADO
        celda.font = LETRA_ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for numero, fila in enumerate(filas, start=2):
        hoja.append(list(fila))
        if numero % 2 == 0:
            for celda in hoja[numero]:
                celda.fill = FILA_PAR
    hoja.freeze_panes = "A2"
    for indice, columna in enumerate(columnas, start=1):
        letra = get_column_letter(indice)
        hoja.column_dimensions[letra].width = max(12, min(44, len(str(columna)) + 2))
        if columna in plegar:
            hoja.column_dimensions[letra].hidden = True
            hoja.column_dimensions[letra].outlineLevel = 1
    if plegar:
        hoja.sheet_properties.outlinePr.summaryRight = False
    return hoja


def _titulo_visible(columna):
    return columna if columna is not None else ""


def _tabla(registros, campos):
    return [[r.get(c, "") for c in campos] for r in registros]


def escribir_base_depurada(ruta, columnas, activos, abandonos, decisiones,
                           cambios, anomalias, vacios, verificacion, reglas,
                           plegar=()):
    """Genera BASE_DATOS_DEPURADA.xlsx con sus siete hojas."""
    libro = openpyxl.Workbook()
    libro.remove(libro.active)

    _escribir_hoja(libro, "ACTIVOS", columnas, activos, plegar)

    columnas_abandonos = ["codigo", "motivo", "desde", "ultimo_acceso",
                          "senales", "quien_decidio"]
    _escribir_hoja(libro, "ABANDONOS", columnas_abandonos,
                   _tabla(decisiones, columnas_abandonos))

    campos = ["columna", "paso", "que_se_hizo", "factor", "por_que"]
    _escribir_hoja(libro, "CAMBIOS", campos, _tabla(cambios, campos))

    campos = ["estudiante", "columna", "valor", "motivo", "detalle", "accion"]
    _escribir_hoja(libro, "ANOMALIAS", campos, _tabla(anomalias, campos))

    campos = ["columna", "tipo_de_columna", "vacios", "figura_pero_sin_dato",
              "no_figura_en_el_archivo", "sin_clasificar", "que_se_hizo"]
    _escribir_hoja(libro, "VACIOS", campos, _tabla(vacios, campos))

    campos = ["columna_a", "columna_b", "relacion_antes", "relacion_despues",
              "veredicto", "detalle"]
    _escribir_hoja(libro, "VERIFICACION", campos, _tabla(verificacion, campos))

    if reglas:
        campos = list(reglas[0].keys())
        _escribir_hoja(libro, "REGLAS_EVALUACION", campos, _tabla(reglas, campos))
    else:
        _escribir_hoja(libro, "REGLAS_EVALUACION", ["aviso"],
                       [["la base de entrada no traia esta hoja"]])

    libro.save(ruta)
    return ruta


def escribir_equivalencia(ruta, columnas, equivalencia):
    """La tabla que permite volver del codigo al estudiante.

    Es el unico archivo de esta etapa que vuelve a contener nombres. Se guarda
    aparte a proposito: la base depurada puede compartirse, esta no.
    """
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "EQUIVALENCIA"
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.fill = ENCABEZADO
        celda.font = LETRA_ENCABEZADO
    for registro in equivalencia:
        hoja.append([registro.get(c, "") for c in columnas])
    hoja.freeze_panes = "A2"
    for indice, columna in enumerate(columnas, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = \
            max(14, min(46, len(str(columna)) + 2))
    libro.save(ruta)
    return ruta


# ----------------------------------------------------------------- informe
def escribir_informe(ruta, materia, resumen, secciones):
    """INFORME_PROCESAMIENTO.docx: portada y un apartado por paso."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()

    doc.add_heading("Informe de procesamiento de datos", level=0)
    subtitulo = doc.add_paragraph()
    corrida = subtitulo.add_run(f"{materia}   ·   Etapa 2 — Procesar")
    corrida.font.size = Pt(14)
    corrida.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(
        f"Generado el {_dt.date.today().strftime('%d/%m/%Y')}. "
        "Esta etapa depura los datos y no saca ninguna conclusion sobre los "
        "estudiantes: eso corresponde a la etapa de analisis.")

    doc.add_heading("Resumen de control", level=1)
    tabla = doc.add_table(rows=1, cols=2)
    tabla.style = "Light Grid Accent 1"
    encabezado = tabla.rows[0].cells
    encabezado[0].text = "Indicador"
    encabezado[1].text = "Resultado"
    for celda in encabezado:
        for parrafo in celda.paragraphs:
            for corrida in parrafo.runs:
                corrida.bold = True
    for etiqueta, valor in resumen:
        fila = tabla.add_row().cells
        fila[0].text = str(etiqueta)
        fila[1].text = str(valor)

    for titulo, descripcion, campos, filas in secciones:
        doc.add_heading(titulo, level=1)
        if descripcion:
            doc.add_paragraph(descripcion)
        if not filas:
            doc.add_paragraph("Sin registros en este apartado.")
            continue
        tabla = doc.add_table(rows=1, cols=len(campos))
        tabla.style = "Light Grid Accent 1"
        encabezado = tabla.rows[0].cells
        for celda, nombre in zip(encabezado, campos):
            celda.text = str(nombre)
            for parrafo in celda.paragraphs:
                for corrida in parrafo.runs:
                    corrida.bold = True
        for registro in filas[:120]:
            fila = tabla.add_row().cells
            for celda, campo in zip(fila, campos):
                celda.text = _texto(registro.get(campo, ""))[:180]
        if len(filas) > 120:
            doc.add_paragraph(
                f"Se muestran las primeras 120 de {len(filas)} filas. "
                "La lista completa esta en el archivo de Excel.")

    for parrafo in doc.paragraphs:
        for corrida in parrafo.runs:
            if not corrida.font.size:
                corrida.font.size = Pt(11)

    doc.save(ruta)
    return ruta
