"""Instalador y autodiagnostico de la Etapa 2 — Procesar.

Se llama instalar_2.py, con el numero de su etapa: cada etapa del ciclo
trae el suyo y asi nunca se confunden ni se pisan entre ellas.

Se ejecuta UNA vez, despues de copiar la carpeta al espacio de trabajo:

    python3 AGENTES/ucb-procesar-datos/instalar_2.py

Comprueba el entorno, verifica que la habilidad esta completa, la prueba de punta a
punta con una base inventada y deja anunciada la habilidad en el AGENTS.md de la
raiz para que el agente la descubra por su cuenta.

No toca ningun archivo del docente.
"""

import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

AQUI = Path(__file__).resolve().parent          # .../AGENTES/ucb-procesar-datos
CARPETA_AGENTES = AQUI.parent                   # .../AGENTES
RAIZ_ESPACIO = CARPETA_AGENTES.parent           # el espacio de trabajo

VERSION_MINIMA = (3, 8)
OK, FALLO, AVISO = "  [OK]   ", "  [FALLO]", "  [AVISO]"

ANUNCIO = """
## Etapa 2 · Procesar

Cuando el docente pida **depurar, limpiar, anonimizar o procesar** su base
integrada, o mencione la Etapa 2 / PROCESAR, lee
`AGENTES/ucb-procesar-datos/SKILL.md` completo y siguelo al pie de la letra,
respetando sus dos pausas. Los archivos de `scripts/` se ejecutan; no se leen ni se
reescriben.
"""


def titulo(texto):
    print(f"\n{texto}\n" + "-" * len(texto))


# --------------------------------------------------------------------- entorno
def comprobar_python():
    actual = ".".join(str(n) for n in sys.version_info[:3])
    if sys.version_info < VERSION_MINIMA:
        minima = ".".join(str(n) for n in VERSION_MINIMA)
        print(f"{FALLO} Python {actual}: se necesita {minima} o superior.")
        print("          Instalalo desde https://www.python.org/downloads/")
        return False
    print(f"{OK} Python {actual}")
    return True


def comprobar_dependencias():
    faltan, todo_bien = [], True
    try:
        import openpyxl
        print(f"{OK} openpyxl {openpyxl.__version__}")
    except ImportError:
        print(f"{FALLO} falta openpyxl (para leer y escribir Excel)")
        faltan.append("openpyxl")
        todo_bien = False
    try:
        import docx  # noqa: F401
        print(f"{OK} python-docx instalado")
    except ImportError:
        print(f"{FALLO} falta python-docx (para el informe en Word)")
        faltan.append("python-docx")
        todo_bien = False
    if faltan:
        print(f"          Instalalas con:  "
              f"{Path(sys.executable).name} -m pip install {' '.join(faltan)}")
    return todo_bien


def comprobar_archivos():
    esperados = [
        AQUI / "SKILL.md",
        AQUI / "scripts" / "procesar.py",
        AQUI / "scripts" / "config.json",
        AQUI / "referencias" / "formato-salida.md",
        AQUI / "ejemplos" / "crear_base_de_prueba.py",
    ]
    faltan = [p for p in esperados if not p.exists()]
    if faltan:
        print(f"{FALLO} la copia esta incompleta. Faltan:")
        for p in faltan:
            print(f"          {p.relative_to(AQUI)}")
        return False
    pasos = sorted(AQUI.glob("scripts/paso_*.py"))
    print(f"{OK} habilidad completa ({len(pasos)} pasos + punto de entrada)")
    return True


# ------------------------------------------------------------- autodiagnostico
def prueba_de_extremo_a_extremo():
    """Corre la habilidad entera sobre una base inventada."""
    with tempfile.TemporaryDirectory(prefix="ucb_e2_") as temporal:
        banco = Path(temporal)
        generador = AQUI / "ejemplos" / "crear_base_de_prueba.py"
        hecho = subprocess.run([sys.executable, str(generador), str(banco)],
                               capture_output=True, text=True)
        if hecho.returncode != 0:
            print(f"{FALLO} no se pudo crear la base de prueba:")
            print(hecho.stderr.strip()[:400])
            return False

        procesar = AQUI / "scripts" / "procesar.py"

        # --- 1. Modo propuesta: no debe escribir nada -----------------------
        hecho = subprocess.run(
            [sys.executable, str(procesar), "CURSO-PRUEBA", "--raiz", str(banco),
             "--proponer"], capture_output=True, text=True)
        if hecho.returncode != 0:
            print(f"{FALLO} el modo propuesta no llego al final:")
            print(hecho.stderr.strip()[:400])
            return False
        bases = banco / "CURSO-PRUEBA-REV" / "bases_de_datos"
        if (bases / "BASE_DATOS_DEPURADA.xlsx").exists():
            print(f"{FALLO} el modo propuesta escribio archivos y no debia")
            return False
        print(f"{OK} el modo propuesta no toca nada")

        # --- 2. Ejecucion real ----------------------------------------------
        hecho = subprocess.run(
            [sys.executable, str(procesar), "CURSO-PRUEBA", "--raiz", str(banco)],
            capture_output=True, text=True)
        if hecho.returncode != 0:
            print(f"{FALLO} la habilidad no llego al final:")
            print(hecho.stderr.strip()[:600])
            return False
        salida = hecho.stdout

        depurada = bases / "BASE_DATOS_DEPURADA.xlsx"
        equivalencia = bases / "EQUIVALENCIA_CODIGOS.xlsx"
        if not depurada.exists():
            print(f"{FALLO} no se genero BASE_DATOS_DEPURADA.xlsx")
            return False

        import openpyxl
        libro = openpyxl.load_workbook(depurada)
        esperadas = ["ACTIVOS", "ABANDONOS", "CAMBIOS", "ANOMALIAS", "VACIOS",
                     "VERIFICACION", "REGLAS_EVALUACION"]
        faltan = [h for h in esperadas if h not in libro.sheetnames]
        if faltan:
            print(f"{FALLO} a la base depurada le faltan hojas: {faltan}")
            libro.close()
            return False
        activos = libro["ACTIVOS"].max_row - 1
        libro.close()
        print(f"{OK} BASE_DATOS_DEPURADA.xlsx con sus 7 hojas ({activos} activos)")

        if not equivalencia.exists():
            print(f"{AVISO} no se genero la tabla de equivalencia de codigos")
        else:
            print(f"{OK} tabla de equivalencia codigo-estudiante generada")

        informe = (banco / "CURSO-PRUEBA-REV" / "RESULTADOS" /
                   "ETAPA 2 - PROCESAR" / "INFORME_PROCESAMIENTO.docx")
        if informe.exists():
            print(f"{OK} informe en Word generado")
        else:
            print(f"{AVISO} no se genero el informe en Word")

        # --- 3. Las garantias que no pueden fallar --------------------------
        if "La cuenta cierra" in salida and "SI" in _valor(salida, "La cuenta cierra"):
            print(f"{OK} la cuenta cierra: activos + abandonos = total de entrada")
        else:
            print(f"{FALLO} la cuenta NO cierra: se perdio alguna fila")
            return False

        if "SI" in _valor(salida, "Archivo de entrada sin modificar"):
            print(f"{OK} la base de la Etapa 1 quedo intacta")
        else:
            print(f"{AVISO} no se pudo confirmar que la entrada quedo intacta")

        huella = _valor(salida, "Huella de la hoja ACTIVOS").strip()
        if huella:
            print(f"{OK} huella de verificacion generada ({huella})")
        return True


def _valor(salida, etiqueta):
    """Lee un valor del resumen de control impreso por el programa."""
    for linea in salida.splitlines():
        if linea.strip().startswith(etiqueta):
            return linea.split(etiqueta, 1)[1]
    return ""


# ------------------------------------------------------------------ AGENTS.md
MARCA = "ucb-procesar-datos"


def anunciar_habilidad():
    destino = RAIZ_ESPACIO / "AGENTS.md"

    if not destino.exists():
        destino.write_text(
            "# Instrucciones para el agente\n\n"
            "Este espacio de trabajo contiene habilidades en `AGENTES/`. Cada una "
            "es una carpeta con un `SKILL.md` que describe un procedimiento, y "
            "programas deterministas que hacen el trabajo pesado.\n"
            + ANUNCIO, encoding="utf-8")
        print(f"{OK} AGENTS.md creado en {destino}")
        return

    actual = destino.read_text(encoding="utf-8", errors="ignore")
    if MARCA in actual:
        print(f"{OK} AGENTS.md ya menciona esta habilidad; no se toco")
        return

    with open(destino, "a", encoding="utf-8") as f:
        f.write("\n" + ANUNCIO)
    print(f"{OK} ya existia un AGENTS.md: se le AGREGO el anuncio de la Etapa 2.")
    print("          No se borro ni se cambio nada de lo que ya tenia.")


# -------------------------------------------------------------------------- main
def main():
    print("=" * 68)
    print(" INSTALACION — Etapa 2: Procesar")
    print("=" * 68)
    print(f"\nEspacio de trabajo detectado: {RAIZ_ESPACIO}")

    titulo("1. Entorno")
    if not (comprobar_python() and comprobar_dependencias()):
        print("\nCorrige lo anterior y vuelve a ejecutar este instalador.")
        sys.exit(1)

    titulo("2. Integridad de la copia")
    if not comprobar_archivos():
        print("\nVuelve a copiar la carpeta ucb-procesar-datos completa.")
        sys.exit(1)

    titulo("3. Prueba con una base inventada")
    print("  (no se usa ningun dato real; la prueba se borra sola al terminar)")
    if not prueba_de_extremo_a_extremo():
        print("\nLa habilidad no paso la prueba. Revisa el mensaje de arriba.")
        sys.exit(1)

    titulo("4. Descubrimiento por el agente")
    anunciar_habilidad()

    print("\n" + "=" * 68)
    print(" LISTO")
    print("=" * 68)
    print("""
Esta etapa necesita el resultado de la Etapa 1. Comprueba que existe
NOMBRE_CARPETA-REV/bases_de_datos/BASE_INTEGRADA.xlsx antes de empezar.

Despues, escribele al agente en tus palabras:

     Quiero procesar la base integrada de mi materia.

Te hara dos preguntas y te mostrara cuatro tablas con todo lo que propone
tocar. Nada se transforma hasta que apruebes.
""")


if __name__ == "__main__":
    main()
