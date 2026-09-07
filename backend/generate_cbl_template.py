"""
generate_cbl_template.py
Genera la Plantilla Oficial de Planificación por Competencias Modular (CBL)
con marcadores unívocos [[TAG]] para inyección automática.
"""
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Aplica color de fondo hexadecimal a una celda de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Establece márgenes internos de una celda en dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_cbl_template(output_path: str):
    doc = Document()
    
    # Configuración de márgenes estándar (0.8 pulgadas para máxima legibilidad técnica)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        # Encabezado formal de calidad institucional
        header = section.header
        p_head = header.paragraphs[0]
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_h = p_head.add_run("PLAN DE APRENDIZAJE MODULAR (CBL) • [[CODIGO_MODULO]] • ISO 21001:2018")
        r_h.font.name = "Arial"
        r_h.font.size = Pt(7.5)
        r_h.font.color.rgb = RGBColor(100, 116, 139)

        # Pie de página institucional
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_f = p_foot.add_run("Carpeta Pedagógica 2.0 • Sistema de Calidad Curricular • DOC-DIR-CP2-PDC-[[CODIGO_MODULO]]")
        r_f.font.name = "Arial"
        r_f.font.size = Pt(7.5)
        r_f.font.color.rgb = RGBColor(100, 116, 139)

    # Colores Atelier / Institucionales
    C_NAVY = RGBColor(26, 58, 92)       # #1A3A5C
    C_SLATE = RGBColor(30, 41, 59)      # #1E293B
    C_AMBER = RGBColor(180, 83, 9)      # #B45309
    C_DARK = RGBColor(15, 23, 42)       # #0F172A
    C_MUTED = RGBColor(100, 116, 139)   # #64748B
    
    # Encabezado Principal y Portada Institucional
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = title_p.add_run("SISTEMA DE GESTIÓN CURRICULAR • EDUCACIÓN SUPERIOR Y FORMACIÓN CIENTÍFICA\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(8.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = C_AMBER
    
    r_title = title_p.add_run("PLAN DE APRENDIZAJE MODULAR POR COMPETENCIAS (CBL)")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = C_NAVY

    p_norma = doc.add_paragraph()
    p_norma.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_norm = p_norma.add_run("Diseño Curricular Inverso (UbD) • Evidencias Auténticas • Enfoque Inclusivo DUA • ISO 21001:2018")
    r_norm.font.name = "Arial"
    r_norm.font.size = Pt(9)
    r_norm.font.italic = True
    r_norm.font.color.rgb = C_MUTED
    
    # 0. Tabla de Control Documental ISO 21001
    t_iso = doc.add_table(rows=4, cols=2)
    t_iso.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_iso.autofit = False

    iso_meta = [
        ("Código Institucional:", "DOC-DIR-CP2-PDC-[[CODIGO_MODULO]]"),
        ("Versión Documental Vigente:", "1.0 (Oficial Aprobado)"),
        ("Marco Normativo / Sistema:", "ISO 21001:2018 (EOMS) Cláusula 7.5 • Información Documentada"),
        ("Vigencia Académica:", "[[PERIODO_MODULO]] • Semestre Académico Oficial")
    ]

    for idx, (lbl, val) in enumerate(iso_meta):
        c0 = t_iso.cell(idx, 0)
        c1 = t_iso.cell(idx, 1)
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(lbl)
        r0.font.name = "Arial"
        r0.font.size = Pt(8.5)
        r0.font.bold = True
        r0.font.color.rgb = C_SLATE

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(8.5)
        r1.font.bold = (idx == 0)
        r1.font.color.rgb = C_NAVY if idx == 0 else C_DARK

    doc.add_paragraph() # Espacio
    
    # 1. Tabla de Datos Generales y Distribución Horaria / Créditos
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("1. DATOS GENERALES DEL MÓDULO Y DISTRIBUCIÓN HORARIA")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(11)
    r_h1.font.bold = True
    r_h1.font.color.rgb = C_SLATE
    
    t_gen = doc.add_table(rows=4, cols=4)
    t_gen.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_gen.autofit = False

    datos_gen = [
        [("Código:", True), ("[[CODIGO_MODULO]]", False), ("Asignatura / Módulo:", True), ("[[NOMBRE_MODULO]]", False)],
        [("Docente Titular:", True), ("[[DOCENTE]]", False), ("Periodo / Módulo:", True), ("[[PERIODO_MODULO]]", False)],
        [("Modalidad:", True), ("[[MODALIDAD]]", False), ("Carga Horaria Total:", True), ("[[DURACION_HORAS]]", False)],
        [("Horas Presenciales/Sincrónicas:", True), ("[[HORAS_PRESENCIALES]]", False), ("Horas Autónomas / Créditos:", True), ("[[HORAS_AUTONOMAS]] • [[CREDITOS_SCT]]", False)]
    ]

    for row_idx, row_data in enumerate(datos_gen):
        for col_idx, (text, is_label) in enumerate(row_data):
            cell = t_gen.cell(row_idx, col_idx)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if is_label:
                r.font.bold = True
                r.font.color.rgb = C_SLATE
                set_cell_background(cell, "F1F5F9")
            else:
                r.font.color.rgb = C_DARK
                set_cell_background(cell, "FFFFFF")

    doc.add_paragraph() # Espacio

    # 2. Competencia Global y Problema del Contexto
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("2. COMPETENCIA GLOBAL Y PROBLEMATIZACIÓN DEL CONTEXTO")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(11)
    r_h2.font.bold = True
    r_h2.font.color.rgb = C_SLATE
    
    t_comp = doc.add_table(rows=2, cols=2)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_comp.autofit = False
    
    c_labels = ["Problema del Contexto Real:", "Competencia Global del Módulo:"]
    c_tags = ["[[PROBLEMA_CONTEXTO]]", "[[COMPETENCIA_GLOBAL]]"]
    
    for idx in range(2):
        cell_lbl = t_comp.cell(idx, 0)
        cell_val = t_comp.cell(idx, 1)
        set_cell_margins(cell_lbl, 100, 100, 120, 120)
        set_cell_margins(cell_val, 100, 100, 120, 120)
        set_cell_background(cell_lbl, "F8FAFC")
        set_cell_background(cell_val, "FFFFFF")
        
        p_l = cell_lbl.paragraphs[0]
        r_l = p_l.add_run(c_labels[idx])
        r_l.font.name = "Arial"
        r_l.font.bold = True
        r_l.font.size = Pt(9)
        r_l.font.color.rgb = C_SLATE
        
        p_v = cell_val.paragraphs[0]
        r_v = p_v.add_run(c_tags[idx])
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9)
        r_v.font.color.rgb = C_DARK

    doc.add_paragraph() # Espacio

    # 3. Matriz de Fases y Desarrollo Curricular Semanal (Hitos 1 a 4)
    h3 = doc.add_paragraph()
    r_h3 = h3.add_run("3. DESARROLLO MODULAR POR FASES Y HITOS DE APRENDIZAJE")
    r_h3.font.name = "Arial"
    r_h3.font.size = Pt(11)
    r_h3.font.bold = True
    r_h3.font.color.rgb = C_SLATE
    
    t_fases = doc.add_table(rows=5, cols=6)
    t_fases.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_fases.autofit = False
    
    # Encabezados de la tabla de fases
    headers = [
        "Fase / Semana", 
        "Saber (Conceptual)", 
        "Hacer (Procedimental)", 
        "Ser y Convivir (Actitudinal)", 
        "Sudor Intelectual (Evidencia)", 
        "Criterio de Evaluación"
    ]
    for col_idx, h_text in enumerate(headers):
        cell = t_fases.cell(0, col_idx)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 100, 100, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for sem in range(1, 5):
        row_cells = t_fases.rows[sem].cells
        bg_col = "FFFFFF" if sem % 2 != 0 else "F8FAFC"
        
        tags_row = [
            f"Semana {sem}:\n[[FASE_{sem}_TITULO]]",
            f"[[FASE_{sem}_SABER]]",
            f"[[FASE_{sem}_HACER]]",
            f"[[FASE_{sem}_SER]]",
            f"[[FASE_{sem}_SUDOR_INTELECTUAL]]",
            f"[[FASE_{sem}_CRITERIO_EVAL]]"
        ]
        
        for col_idx, tag_text in enumerate(tags_row):
            cell = row_cells[col_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, 90, 90, 80, 80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(tag_text)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = C_DARK
            if col_idx == 0:
                r.font.bold = True

    doc.add_paragraph() # Espacio

    # 4. Inclusión DUA y Triangulación Anti-Outsourcing
    h4 = doc.add_paragraph()
    r_h4 = h4.add_run("4. ENFOQUE HUMANO: INCLUSIÓN (DUA) Y RIGOR EPISTÉMICO")
    r_h4.font.name = "Arial"
    r_h4.font.size = Pt(11)
    r_h4.font.bold = True
    r_h4.font.color.rgb = C_SLATE
    
    t_dua = doc.add_table(rows=2, cols=2)
    t_dua.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_dua.autofit = False
    
    dua_labels = [
        "Diseño Universal para el Aprendizaje (DUA CAST 2024):",
        "Estrategia Anti-Outsourcing y Triangulación Socrática:"
    ]
    dua_tags = [
        "[[ADAPTACIONES_DUA_INCLUSION]]",
        "[[ESTRATEGIA_ANTI_OUTSOURCING]]"
    ]
    
    for idx in range(2):
        cell_lbl = t_dua.cell(idx, 0)
        cell_val = t_dua.cell(idx, 1)
        set_cell_margins(cell_lbl, 100, 100, 120, 120)
        set_cell_margins(cell_val, 100, 100, 120, 120)
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "FFFFFF")
        
        p_l = cell_lbl.paragraphs[0]
        r_l = p_l.add_run(dua_labels[idx])
        r_l.font.name = "Arial"
        r_l.font.bold = True
        r_l.font.size = Pt(9)
        r_l.font.color.rgb = C_SLATE
        
        p_v = cell_val.paragraphs[0]
        r_v = p_v.add_run(dua_tags[idx])
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9)
        r_v.font.color.rgb = C_DARK

    doc.add_paragraph() # Espacio

    # 5. Producto Final Integrador
    h5 = doc.add_paragraph()
    r_h5 = h5.add_run("5. PRODUCTO FINAL INTEGRADOR DEL MÓDULO")
    r_h5.font.name = "Arial"
    r_h5.font.size = Pt(11)
    r_h5.font.bold = True
    r_h5.font.color.rgb = C_SLATE
    
    t_prod = doc.add_table(rows=1, cols=2)
    t_prod.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_prod.autofit = False
    
    c_p_lbl = t_prod.cell(0, 0)
    c_p_val = t_prod.cell(0, 1)
    set_cell_margins(c_p_lbl, 100, 100, 120, 120)
    set_cell_margins(c_p_val, 100, 100, 120, 120)
    set_cell_background(c_p_lbl, "FEF3C7") # Ámbar suave
    set_cell_background(c_p_val, "FFFFFF")
    
    p_pl = c_p_lbl.paragraphs[0]
    r_pl = p_pl.add_run("Evidencia Auténtica Integradora:")
    r_pl.font.name = "Arial"
    r_pl.font.bold = True
    r_pl.font.size = Pt(9)
    r_pl.font.color.rgb = C_AMBER
    
    p_pv = c_p_val.paragraphs[0]
    r_pv = p_pv.add_run("[[PRODUCTO_FINAL_MODULAR]]")
    r_pv.font.name = "Arial"
    r_pv.font.size = Pt(9)
    r_pv.font.color.rgb = C_DARK

    doc.add_paragraph() # Espacio

    # 6. Ciclo de Firmas de Calidad Institucional (Tripartita)
    h6 = doc.add_paragraph()
    r_h6 = h6.add_run("6. FIRMAS DE APROBACIÓN Y ASEGURAMIENTO DE LA CALIDAD (ISO 21001)")
    r_h6.font.name = "Arial"
    r_h6.font.size = Pt(11)
    r_h6.font.bold = True
    r_h6.font.color.rgb = C_SLATE

    t_sign = doc.add_table(rows=2, cols=3)
    t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sign.autofit = False

    sign_roles = [
        ("DOCENTE TITULAR", "Elaboración y custodia de evidencias", "[[DOCENTE]]"),
        ("COORDINACIÓN DE CARRERA", "Revisión curricular y pertinencia", "Comisión Curricular de Área"),
        ("DIRECCIÓN ACADÉMICA / DECANATO", "Aprobación oficial institucional", "Dirección de Calidad Académica")
    ]

    for col_idx, (role, desc, name) in enumerate(sign_roles):
        c_top = t_sign.cell(0, col_idx)
        c_bot = t_sign.cell(1, col_idx)
        set_cell_margins(c_top, 300, 60, 60, 60) # Espacio para firma
        set_cell_margins(c_bot, 40, 80, 60, 60)
        set_cell_background(c_top, "FAFAFA")
        set_cell_background(c_bot, "F1F5F9")

        p_t = c_top.paragraphs[0]
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_line = p_t.add_run("____________________________\nFirma y Sello")
        r_line.font.name = "Arial"
        r_line.font.size = Pt(8)
        r_line.font.color.rgb = C_MUTED

        p_b = c_bot.paragraphs[0]
        p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_b1 = p_b.add_run(f"{role}\n")
        r_b1.font.name = "Arial"
        r_b1.font.bold = True
        r_b1.font.size = Pt(8)
        r_b1.font.color.rgb = C_SLATE

        r_b2 = p_b.add_run(f"{name}\n({desc})")
        r_b2.font.name = "Arial"
        r_b2.font.size = Pt(7)
        r_b2.font.color.rgb = C_MUTED

    # Pie final editorial
    doc.add_paragraph()
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_foot = p_foot.add_run("Generado automáticamente por el Motor de Planificación Modular CBL • Carpeta Pedagógica 2.0 • ISO 21001:2018")
    r_foot.font.name = "Arial"
    r_foot.font.size = Pt(7.5)
    r_foot.font.italic = True
    r_foot.font.color.rgb = C_MUTED
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Plantilla CBL guardada con éxito en: {output_path}")

if __name__ == "__main__":
    t_dir = os.path.join(os.path.dirname(__file__), "templates")
    out = os.path.join(t_dir, "Plantilla_CBL_Modular.docx")
    create_cbl_template(out)
