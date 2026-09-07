"""
research_service.py - Servicio de Investigacion Comparativa
Carpeta Pedagogica 2.0 | Luis Alfredo Andia Valverde
Licencia CC BY-NC 4.0
"""

import os
import json
import statistics
import re
from typing import Optional

BASE_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "uploads")


def _sanitize(name: str) -> str:
    clean = name.replace("..", "").replace("/", " ").replace("\\", " ")
    clean = re.sub(r'[<>:"/\\|?*]', '', clean)
    return re.sub(r'\s+', ' ', clean).strip()


def _resolve_rev_dir(subject_name: str) -> Optional[str]:
    clean = _sanitize(subject_name)
    if clean.endswith("-REV"):
        clean = clean[:-4].strip()
    if not os.path.exists(BASE_DIR):
        return None
    clean_lower = clean.lower()
    for entry in os.listdir(BASE_DIR):
        if not entry.endswith("-REV"):
            continue
        base = entry[:-4].strip()
        if base.lower() == clean_lower or clean_lower in base.lower() or base.lower().startswith(clean_lower):
            path = os.path.join(BASE_DIR, entry)
            if os.path.isdir(path):
                return path
    return None


def get_subjects_index() -> dict:
    if not os.path.exists(BASE_DIR):
        return {"by_group": {}, "by_code": {}, "all": []}
    by_group: dict = {}
    by_code: dict = {}
    all_subjects = []
    for entry in os.listdir(BASE_DIR):
        if not entry.endswith("-REV"):
            continue
        rev_path = os.path.join(BASE_DIR, entry)
        manifest = {}
        mp = os.path.join(rev_path, "manifiesto.json")
        if os.path.exists(mp):
            try:
                with open(mp, "r", encoding="utf-8") as f:
                    manifest = json.load(f)
            except Exception:
                pass
        name = entry[:-4].strip()
        code = manifest.get("code", "")
        group = manifest.get("group", "")
        asignatura = manifest.get("asignatura", manifest.get("name", name))
        year = manifest.get("year", "")
        period = manifest.get("period", "")

        # Inferencia inteligente de codigo y asignatura si no estan en manifiesto
        if not code and " - " in name:
            parts = name.split(" - ", 1)
            if len(parts[0]) <= 10 and any(ch.isdigit() for ch in parts[0]):
                code = parts[0].strip()
                if not manifest.get("asignatura"):
                    asignatura = parts[1].strip()

        # Inferencia de grupo si viene en parentesis ej: Didactica (Grupo 2026)
        if not group and "(" in name and name.endswith(")"):
            try:
                base_part, grp_part = name.rsplit("(", 1)
                group = grp_part.rstrip(")").strip()
            except Exception:
                pass

        # Fallback de grupo para asignaturas existentes sin grupo explicito
        display_group = group if group else "Cohorte General"

        record = {
            "name": name, "code": code, "asignatura": asignatura,
            "group": group, "year": year or 2026, "period": period or "1",
            "has_insights": os.path.exists(os.path.join(rev_path, "insights.json")),
            "has_base_limpia": os.path.exists(os.path.join(rev_path, "bases_de_datos", "BASE_LIMPIA_ANONIMIZADA.xlsx")),
        }
        all_subjects.append(record)

        # Agrupar por grupo (Eje A)
        grp_key = f"{display_group} ({year or 2026})"
        by_group.setdefault(grp_key, []).append(record)

        # Agrupar por codigo (Eje B)
        if code:
            by_code.setdefault(code, {"code": code, "asignatura": asignatura, "groups": []})["groups"].append(record)
    return {"by_group": by_group, "by_code": by_code, "all": all_subjects}


def _safe_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None

def _safe_int(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None

def _find_col(headers, candidates):
    hl = [h.lower() for h in headers]
    for c in candidates:
        if c.lower() in hl:
            return hl.index(c.lower())
    return None

def _pearson(x, y):
    n = len(x)
    mx, my = sum(x)/n, sum(y)/n
    num = sum((xi-mx)*(yi-my) for xi, yi in zip(x, y))
    den = (sum((xi-mx)**2 for xi in x) * sum((yi-my)**2 for yi in y)) ** 0.5
    return (num / den) if den != 0 else 0.0


def load_subject_metrics(subject_name: str) -> dict:
    rev_dir = _resolve_rev_dir(subject_name)
    if not rev_dir:
        return {"error": f"No encontrado: {subject_name}", "name": subject_name}
    manifest = {}
    mp = os.path.join(rev_dir, "manifiesto.json")
    if os.path.exists(mp):
        try:
            with open(mp, "r", encoding="utf-8") as f:
                manifest = json.load(f)
        except Exception:
            pass
    display_name = os.path.basename(rev_dir)[:-4].strip()
    code = manifest.get("code", "")
    group = manifest.get("group", "")
    asignatura = manifest.get("asignatura", manifest.get("name", subject_name))
    year = manifest.get("year", "")
    period = manifest.get("period", "")

    if not code and " - " in display_name:
        parts = display_name.split(" - ", 1)
        if len(parts[0]) <= 10 and any(ch.isdigit() for ch in parts[0]):
            code = parts[0].strip()
            if not manifest.get("asignatura"):
                asignatura = parts[1].strip()

    if not group and "(" in display_name and display_name.endswith(")"):
        try:
            base_part, grp_part = display_name.rsplit("(", 1)
            group = grp_part.rstrip(")").strip()
        except Exception:
            pass

    metrics = {
        "name": subject_name,
        "display_name": display_name,
        "code": code,
        "asignatura": asignatura,
        "group": group,
        "year": year or 2026,
        "period": period or "1",
        "promedio": None, "desviacion_estandar": None, "mediana": None,
        "minimo": None, "maximo": None, "pct_aprobados": None,
        "n_riesgo_abandono": None, "n_estudiantes": None,
        "pct_asistencia": None, "sudor_intelectual_promedio": None,
        "nivel_autonomia": None, "zdp": None, "perfil_general": None,
        "data_source": "sin_datos",
    }
    # Fuente 1: insights.json
    ip = os.path.join(rev_dir, "insights.json")
    if os.path.exists(ip):
        try:
            with open(ip, "r", encoding="utf-8") as f:
                ins = json.load(f)
            metrics["promedio"] = _safe_float(ins.get("promedio_general") or ins.get("promedio"))
            metrics["desviacion_estandar"] = _safe_float(ins.get("desviacion_estandar") or ins.get("std"))
            metrics["n_riesgo_abandono"] = _safe_int(ins.get("riesgo_abandono") or ins.get("n_riesgo"))
            metrics["nivel_autonomia"] = ins.get("nivel_autonomia") or ins.get("autonomia")
            metrics["zdp"] = ins.get("zdp_dominante") or ins.get("zdp")
            metrics["perfil_general"] = ins.get("perfil_general") or ins.get("perfil")
            if ins.get("pct_aprobados"):
                metrics["pct_aprobados"] = _safe_float(ins.get("pct_aprobados"))
            dist = ins.get("distribucion_notas") or ins.get("distribucion") or {}
            if dist and metrics["pct_aprobados"] is None:
                apr = sum(v for k, v in dist.items() if _safe_float(k) is not None and _safe_float(k) >= 60)
                tot = sum(dist.values())
                if tot > 0:
                    metrics["pct_aprobados"] = round(apr / tot * 100, 1)
            metrics["data_source"] = "insights_json"
        except Exception:
            pass
    # Fuente 2: BASE_LIMPIA_ANONIMIZADA.xlsx
    if metrics["promedio"] is None:
        for fname in ["bases_de_datos/BASE_LIMPIA_ANONIMIZADA.xlsx", "bases_de_datos/BASE_INTEGRADA.xlsx"]:
            bp = os.path.join(rev_dir, fname)
            if os.path.exists(bp):
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(bp, read_only=True, data_only=True)
                    ws = wb.active
                    headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
                    pc = _find_col(headers, ["Promedio_Ponderado","Promedio","Nota Final","Nota_Final","Puntaje_Total","Calificacion"])
                    ac = _find_col(headers, ["Porcentaje_%","Asistencia_%","Asistencia"])
                    sc = _find_col(headers, ["Iteraciones_Prompt_Total","Iteraciones","Sudor_Intelectual"])
                    notas, avls, svls = [], [], []
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if pc is not None and row[pc] is not None:
                            v = _safe_float(row[pc])
                            if v is not None: notas.append(v)
                        if ac is not None and row[ac] is not None:
                            v = _safe_float(row[ac])
                            if v is not None: avls.append(v)
                        if sc is not None and row[sc] is not None:
                            v = _safe_float(row[sc])
                            if v is not None: svls.append(v)
                    wb.close()
                    if notas:
                        metrics["promedio"] = round(statistics.mean(notas), 2)
                        metrics["desviacion_estandar"] = round(statistics.stdev(notas) if len(notas)>1 else 0, 2)
                        metrics["mediana"] = round(statistics.median(notas), 2)
                        metrics["minimo"] = round(min(notas), 2)
                        metrics["maximo"] = round(max(notas), 2)
                        metrics["n_estudiantes"] = len(notas)
                        metrics["pct_aprobados"] = round(sum(1 for n in notas if n>=60)/len(notas)*100, 1)
                        metrics["n_riesgo_abandono"] = sum(1 for n in notas if n<55)
                        metrics["data_source"] = "xlsx_directo"
                    if avls: metrics["pct_asistencia"] = round(statistics.mean(avls), 1)
                    if svls: metrics["sudor_intelectual_promedio"] = round(statistics.mean(svls), 1)
                    break
                except Exception:
                    pass
    # Fuente 3: notas originales raw
    if metrics["promedio"] is None:
        raw_dir = os.path.join(BASE_DIR, os.path.basename(rev_dir)[:-4].strip())
        fp = os.path.join(raw_dir, "notas_evaluaciones.xlsx")
        if os.path.exists(fp):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
                ws = wb.active
                headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
                pc = _find_col(headers, ["Promedio_Ponderado","Promedio","Nota Final","Nota_Final"])
                notas = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if pc is not None and row[pc] is not None:
                        v = _safe_float(row[pc])
                        if v is not None: notas.append(v)
                wb.close()
                if notas:
                    metrics["promedio"] = round(statistics.mean(notas), 2)
                    metrics["desviacion_estandar"] = round(statistics.stdev(notas) if len(notas)>1 else 0, 2)
                    metrics["mediana"] = round(statistics.median(notas), 2)
                    metrics["minimo"] = round(min(notas), 2)
                    metrics["maximo"] = round(max(notas), 2)
                    metrics["n_estudiantes"] = len(notas)
                    metrics["pct_aprobados"] = round(sum(1 for n in notas if n>=60)/len(notas)*100, 1)
                    metrics["n_riesgo_abandono"] = sum(1 for n in notas if n<55)
                    metrics["data_source"] = "xlsx_raw"
            except Exception:
                pass
    return metrics


def compare_subjects(subject_names: list, focus_metrics: list = None) -> dict:
    all_metrics, errors = [], []
    for name in subject_names:
        m = load_subject_metrics(name)
        if "error" in m:
            errors.append(m["error"])
        else:
            all_metrics.append(m)
    if not all_metrics:
        return {"error": "No se pudieron cargar metricas", "errors": errors}
    quant_keys = ["promedio","desviacion_estandar","pct_aprobados","n_riesgo_abandono","pct_asistencia","sudor_intelectual_promedio"]
    active_keys = [k for k in (focus_metrics or quant_keys) if k in quant_keys]
    global_stats = {}
    for key in active_keys:
        values = [m[key] for m in all_metrics if m.get(key) is not None]
        if values:
            mv = statistics.mean(values)
            sv = statistics.stdev(values) if len(values)>1 else 0
            global_stats[key] = {"mean": round(mv,2),"std": round(sv,2),"min": round(min(values),2),"max": round(max(values),2)}
    for m in all_metrics:
        m["deltas"] = {}; m["z_scores"] = {}; m["rankings"] = {}
        for key in active_keys:
            if m.get(key) is not None and key in global_stats:
                gs = global_stats[key]
                m["deltas"][key] = round(m[key]-gs["mean"], 2)
                m["z_scores"][key] = round((m[key]-gs["mean"])/gs["std"], 2) if gs["std"]>0 else 0.0
    for key in active_keys:
        available = [(m["name"], m[key]) for m in all_metrics if m.get(key) is not None]
        reverse = key not in ["n_riesgo_abandono","desviacion_estandar"]
        available.sort(key=lambda x: x[1], reverse=reverse)
        for rank, (nm, _) in enumerate(available, 1):
            for m in all_metrics:
                if m["name"] == nm:
                    m["rankings"][key] = rank
    correlations = {}
    if len(all_metrics) >= 3:
        pairs = [("promedio","pct_asistencia"),("promedio","sudor_intelectual_promedio"),("promedio","pct_aprobados"),("pct_asistencia","n_riesgo_abandono")]
        for k1,k2 in pairs:
            v1 = [m[k1] for m in all_metrics if m.get(k1) is not None and m.get(k2) is not None]
            v2 = [m[k2] for m in all_metrics if m.get(k1) is not None and m.get(k2) is not None]
            if len(v1)>=3:
                try: correlations[f"{k1}_vs_{k2}"] = round(_pearson(v1,v2),3)
                except Exception: pass
    srt = sorted([m for m in all_metrics if m.get("promedio") is not None], key=lambda x: x["promedio"], reverse=True)
    return {
        "subjects": all_metrics, "global_stats": global_stats, "active_metrics": active_keys,
        "correlations": correlations,
        "best_overall": srt[0]["name"] if srt else None,
        "worst_overall": srt[-1]["name"] if srt else None,
        "n_compared": len(all_metrics), "errors": errors,
    }


RESEARCH_LINES = [
    {"id":"rendimiento_asignaturas","label":"Rendimiento diferencial entre asignaturas","description":"En que asignaturas rinde mejor/peor un mismo grupo y por que?","eje":"A","hipotesis_ejemplo":"Los estudiantes del Grupo 2026 presentan rendimiento inferior en ciencias exactas respecto a humanidades.","metricas_clave":["promedio","desviacion_estandar","pct_aprobados","n_riesgo_abandono"]},
    {"id":"efecto_docente_cohorte","label":"Efecto docente/metodologia entre cohortes","description":"Cambia el rendimiento de la misma asignatura entre grupos sucesivos?","eje":"B","hipotesis_ejemplo":"La metodologia CBL produjo aumento significativo en aprobacion entre grupos sucesivos.","metricas_clave":["promedio","pct_aprobados","sudor_intelectual_promedio","n_riesgo_abandono"]},
    {"id":"engagement_vs_rendimiento","label":"Engagement (sudor intelectual) vs rendimiento","description":"Los grupos con mayor esfuerzo cognitivo obtienen mejores resultados?","eje":"A","hipotesis_ejemplo":"Existe correlacion positiva entre Sudor Intelectual y promedio de calificaciones.","metricas_clave":["promedio","sudor_intelectual_promedio","nivel_autonomia","pct_aprobados"]},
    {"id":"asistencia_vs_aprobacion","label":"Asistencia y riesgo de abandono","description":"Que tan predictiva es la asistencia sobre la aprobacion?","eje":"A","hipotesis_ejemplo":"Asistencia < 70% es predictor significativo de reprobacion en asignaturas de alta carga practica.","metricas_clave":["pct_asistencia","pct_aprobados","n_riesgo_abandono","promedio"]},
    {"id":"equidad_varianza","label":"Equidad y dispersion del aprendizaje","description":"Hay asignaturas donde la brecha de rendimiento intragrupo es mayor?","eje":"A","hipotesis_ejemplo":"La varianza en QUIM101 es significativamente mayor que en INV101, evidenciando polarizacion del aprendizaje.","metricas_clave":["desviacion_estandar","promedio","minimo","maximo","pct_aprobados"]},
    {"id":"personalizado","label":"Linea personalizada (con asistencia IA)","description":"Describe tu pregunta de investigacion y Gemini te ayudara a formalizarla.","eje":"A/B","hipotesis_ejemplo":"","metricas_clave":[]},
]

def get_research_lines() -> list:
    return RESEARCH_LINES


def build_hypothesis_prompt(research_line_id: str, context: dict, custom_question: str = "") -> str:
    line = next((l for l in RESEARCH_LINES if l["id"] == research_line_id), None)
    subjects_ctx = json.dumps([{"nombre": s.get("asignatura") or s.get("name"),"codigo": s.get("code"),"grupo": s.get("group"),"anyo": s.get("year")} for s in context.get("selected_subjects", [])], ensure_ascii=False)
    if research_line_id == "personalizado" and custom_question:
        desc = f"Pregunta exploratoria del docente: {custom_question}"
    elif line:
        desc = f"Linea: {line['label']}. {line['description']}"
    else:
        desc = "Investigacion educativa comparativa general."
    return f"""Eres un experto en investigacion educativa cuantitativa con enfoque humanista (Vygotsky, Deci & Ryan, CBL).
El docente quiere realizar una investigacion comparativa entre:
{subjects_ctx}

{desc}

Proporciona en formato JSON estricto:
{{
  "hipotesis_nula": "H0 formal estadistica",
  "hipotesis_alternativa": "H1 formal estadistica",
  "pregunta_investigacion": "Pregunta clara y medible",
  "metricas_recomendadas": ["lista de metricas de: promedio, desviacion_estandar, pct_aprobados, n_riesgo_abandono, pct_asistencia, sudor_intelectual_promedio"],
  "diseno_metodologico": "Breve descripcion: comparativo, longitudinal, correlacional, etc.",
  "advertencias": ["Limitaciones importantes del analisis"],
  "interpretacion_esperada": "Que resultado confirmaria o refutaria la hipotesis"
}}"""


def build_comparison_narrative_prompt(comparison_data: dict, hypothesis: dict, teacher_name: str = "Docente") -> str:
    subjects_summary = [{"nombre": s.get("asignatura") or s.get("display_name"),"grupo": s.get("group"),"promedio": s.get("promedio"),"pct_aprobados": s.get("pct_aprobados"),"n_riesgo": s.get("n_riesgo_abandono"),"desv_std": s.get("desviacion_estandar"),"asistencia": s.get("pct_asistencia"),"delta_promedio": s.get("deltas",{}).get("promedio"),"z_score": s.get("z_scores",{}).get("promedio")} for s in comparison_data.get("subjects",[])]
    return f"""Eres investigador educativo experto. Redacta un informe comparativo formal.
Docente: {teacher_name}
Hipotesis: {hypothesis.get("hipotesis_alternativa","No especificada")}
Pregunta: {hypothesis.get("pregunta_investigacion","No especificada")}
Datos: {json.dumps(subjects_summary, ensure_ascii=False)}
Correlaciones: {json.dumps(comparison_data.get("correlations",{}), ensure_ascii=False)}

Redacta con secciones:
1. Resumen ejecutivo
2. Hallazgos principales (con datos numericos especificos)
3. Interpretacion pedagogica (Vygotsky/Deci & Ryan)
4. Conclusion respecto a la hipotesis
5. Recomendaciones de intervencion docente"""


def export_comparison_xlsx(comparison_data: dict, hypothesis: dict, narrative: str = "") -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    import io
    wb = openpyxl.Workbook()
    # Hoja 1: Resumen
    ws1 = wb.active; ws1.title = "Resumen Investigacion"
    ws1.append(["INVESTIGACION COMPARATIVA - EVIDENCIA OBJETIVA"])
    ws1["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws1["A1"].fill = PatternFill("solid", fgColor="0D1B2A")
    ws1.merge_cells("A1:F1")
    ws1.append([])
    ws1.append(["Hipotesis H1:", hypothesis.get("hipotesis_alternativa","—")])
    ws1.append(["Pregunta:", hypothesis.get("pregunta_investigacion","—")])
    ws1.append(["Diseno:", hypothesis.get("diseno_metodologico","—")])
    ws1.append(["N asignaturas comparadas:", comparison_data.get("n_compared",0)])
    ws1.column_dimensions["A"].width = 28; ws1.column_dimensions["B"].width = 70
    # Hoja 2: Datos Comparativos
    ws2 = wb.create_sheet("Datos Comparativos")
    hdrs = ["Asignatura","Grupo","Anyo","Promedio","Desv Std","% Aprobados","N Riesgo","% Asistencia","Sudor Intelectual","Delta Promedio","Z-Score"]
    ws2.append(hdrs)
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1A3A5C")
        c.alignment = Alignment(horizontal="center")
    for s in comparison_data.get("subjects",[]):
        row = [s.get("asignatura") or s.get("display_name",""),s.get("group",""),s.get("year",""),
               s.get("promedio"),s.get("desviacion_estandar"),s.get("pct_aprobados"),
               s.get("n_riesgo_abandono"),s.get("pct_asistencia"),s.get("sudor_intelectual_promedio"),
               s.get("deltas",{}).get("promedio"),s.get("z_scores",{}).get("promedio")]
        ws2.append(row)
        r = ws2.max_row
        # Semaforo promedio
        v = row[3]
        if v is not None:
            color = "C6EFCE" if v>=80 else "FFEB9C" if v>=60 else "FFC7CE"
            ws2.cell(r,4).fill = PatternFill("solid", fgColor=color)
        # Semaforo aprobados
        v2 = row[5]
        if v2 is not None:
            color = "C6EFCE" if v2>=85 else "FFEB9C" if v2>=70 else "FFC7CE"
            ws2.cell(r,6).fill = PatternFill("solid", fgColor=color)
    for col in ws2.columns:
        ws2.column_dimensions[get_column_letter(col[0].column)].width = 18
    # Hoja 3: Correlaciones
    ws3 = wb.create_sheet("Correlaciones")
    ws3.append(["Par de variables","r de Pearson","Interpretacion"])
    for c in ws3[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1A3A5C")
    for pair, r in comparison_data.get("correlations",{}).items():
        interp = "Fuerte positiva" if r>=0.7 else "Moderada positiva" if r>=0.4 else "Fuerte negativa" if r<=-0.7 else "Moderada negativa" if r<=-0.4 else "Debil o sin correlacion"
        ws3.append([pair.replace("_"," ").replace("vs","vs."), r, interp])
    ws3.column_dimensions["A"].width = 38; ws3.column_dimensions["B"].width = 16; ws3.column_dimensions["C"].width = 28
    # Hoja 4: Global Stats
    ws4 = wb.create_sheet("Estadisticas Globales")
    ws4.append(["Metrica","Media Global","Desv Std","Minimo","Maximo"])
    for c in ws4[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2D6A4F")
    for key, st in comparison_data.get("global_stats",{}).items():
        ws4.append([key.replace("_"," ").title(), st["mean"], st["std"], st["min"], st["max"]])
    # Hoja 5: Narrativo IA
    if narrative:
        ws5 = wb.create_sheet("Informe Narrativo IA")
        ws5.column_dimensions["A"].width = 12; ws5.column_dimensions["B"].width = 90
        ws5.append(["INFORME NARRATIVO - ANALISIS COMPARATIVO CON IA"])
        ws5["A1"].font = Font(bold=True, size=12, color="FFFFFF")
        ws5["A1"].fill = PatternFill("solid", fgColor="0D1B2A")
        ws5.merge_cells("A1:B1")
        ws5.append([])
        for para in narrative.split("\n"):
            if para.strip():
                ws5.append(["", para.strip()])
                ws5[ws5.max_row][1].alignment = Alignment(wrap_text=True)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()


def export_comparison_docx(comparison_data: dict, hypothesis: dict, narrative: str, teacher_name: str = "Docente") -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import io
    from datetime import datetime
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5); section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3); section.right_margin = Cm(2.5)
    # Portada
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("INFORME DE INVESTIGACION EDUCATIVA COMPARATIVA")
    run.font.size = Pt(16); run.font.bold = True; run.font.color.rgb = RGBColor(0x1A,0x3A,0x5C)
    doc.add_paragraph()
    for text in [f"Docente: {teacher_name}", f"Fecha: {datetime.now().strftime('%d/%m/%Y')}", f"Asignaturas comparadas: {comparison_data.get('n_compared',0)}", "Generado con Carpeta Pedagogica 2.0 | ISO 21001:2018"]:
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(text); r2.font.size = Pt(11 if "Docente" in text or "Fecha" in text or "Asig" in text else 9)
    doc.add_page_break()
    # Marco de investigacion
    doc.add_heading("1. Marco de Investigacion", 1)
    doc.add_heading("Pregunta de Investigacion", 2)
    p3 = doc.add_paragraph(); r3 = p3.add_run(hypothesis.get("pregunta_investigacion","No definida")); r3.font.italic = True
    doc.add_heading("Hipotesis", 2)
    doc.add_paragraph(f"H0: {hypothesis.get('hipotesis_nula','No definida')}")
    doc.add_paragraph(f"H1: {hypothesis.get('hipotesis_alternativa','No definida')}")
    doc.add_heading("Diseno Metodologico", 2)
    doc.add_paragraph(hypothesis.get("diseno_metodologico","—"))
    if hypothesis.get("advertencias"):
        doc.add_heading("Limitaciones", 2)
        for adv in hypothesis.get("advertencias",[]):
            p4 = doc.add_paragraph(style="List Bullet"); p4.add_run(adv).font.size = Pt(10)
    # Tabla comparativa
    doc.add_heading("2. Datos Comparativos", 1)
    subjects = comparison_data.get("subjects",[])
    if subjects:
        hdrs = ["Asignatura","Grupo","Promedio","Sigma","% Aprobados","Riesgo","Delta Media"]
        tbl = doc.add_table(rows=1+len(subjects), cols=len(hdrs)); tbl.style = "Table Grid"
        for i,h in enumerate(hdrs):
            cell = tbl.rows[0].cells[i]; cell.text = h
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            from docx.oxml.ns import qn; from docx.oxml import OxmlElement
            tc = cell._tc; tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"),"1A3A5C"); shd.set(qn("w:color"),"auto"); shd.set(qn("w:val"),"clear")
            tcPr.append(shd)
        for ri,s in enumerate(subjects,1):
            row = tbl.rows[ri]
            vals = [s.get("asignatura") or s.get("display_name",""),s.get("group",""),
                    f"{s['promedio']:.1f}" if s.get("promedio") is not None else "N/A",
                    f"{s['desviacion_estandar']:.1f}" if s.get("desviacion_estandar") is not None else "N/A",
                    f"{s['pct_aprobados']:.1f}%" if s.get("pct_aprobados") is not None else "N/A",
                    str(s.get("n_riesgo_abandono","N/A")),
                    f"{s.get('deltas',{}).get('promedio',0):+.1f}" if s.get("deltas") else "N/A"]
            for ci,val in enumerate(vals):
                row.cells[ci].text = str(val)
    doc.add_paragraph()
    # Narrativo
    doc.add_heading("3. Analisis e Interpretacion (IA)", 1)
    if narrative:
        for para in narrative.split("\n"):
            if para.strip():
                if any(para.strip().startswith(f"{n}.") for n in range(1,10)):
                    doc.add_heading(para.strip(), 2)
                elif para.strip().startswith(("-","*","•")):
                    p5 = doc.add_paragraph(style="List Bullet"); p5.add_run(para.strip().lstrip("-*• ")).font.size = Pt(10)
                else:
                    doc.add_paragraph(para.strip())
    # Correlaciones
    corr = comparison_data.get("correlations",{})
    if corr:
        doc.add_heading("4. Correlaciones Estadisticas", 1)
        for pair,r in corr.items():
            label = pair.replace("_"," ").replace("vs","vs.")
            interp = "correlacion fuerte positiva" if r>=0.7 else "correlacion moderada positiva" if r>=0.4 else "correlacion fuerte negativa" if r<=-0.7 else "correlacion moderada negativa" if r<=-0.4 else "correlacion debil"
            p6 = doc.add_paragraph(style="List Bullet"); p6.add_run(f"{label}: r = {r:.3f} ({interp})").font.size = Pt(10)
    buf = io.BytesIO(); doc.save(buf); buf.seek(0)
    return buf.read()
