"""
version_control_service.py
Sistema de Control Documental y Trazabilidad de Versiones (ISO 21001:2018 Cláusula 7.5)
Genera y mantiene sincronizado el libro oficial:
uploads/{materia}-REV/bases_de_datos/CONTROL_VERSIONES_DOCUMENTAL.xlsx
"""
import os
import hashlib
import re
from datetime import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_UPLOADS = os.path.join(os.path.dirname(os.path.dirname(__file__)), "uploads")

def calculate_sha256(file_path: str) -> str:
    """Calcula el hash SHA-256 de un archivo para garantizar la inmutabilidad y detectar alteraciones."""
    if not os.path.exists(file_path):
        return "ARCHIVO_NO_ENCONTRADO"
    sha256 = hashlib.sha256()
    try:
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(65536), b""):
                sha256.update(chunk)
        return sha256.hexdigest()[:16] + "..." # Hash abreviado para visualización ejecutiva
    except Exception:
        return "ERROR_CALCULO_HASH"

def get_version_control_path(subject_name: str) -> str:
    """Devuelve la ruta absoluta del archivo de control de versiones de la materia."""
    clean_subj = re.sub(r'[<>:"/\\|?*]', '', subject_name).strip()
    if clean_subj.endswith("-REV"):
        clean_subj = clean_subj[:-4].strip()
        
    rev_dir = os.path.join(BASE_UPLOADS, f"{clean_subj}-REV")
    db_dir = os.path.join(rev_dir, "bases_de_datos")
    os.makedirs(db_dir, exist_ok=True)
    return os.path.join(db_dir, "CONTROL_VERSIONES_DOCUMENTAL.xlsx")

def init_version_control_workbook(subject_name: str) -> str:
    """
    Crea la estructura del libro Excel de Control de Versiones con 3 hojas institucionales:
    1. Matriz_Control_Documentos
    2. Historial_Trazabilidad_ISO
    3. Normativa_y_Advertencias
    """
    wb_path = get_version_control_path(subject_name)
    clean_subj = re.sub(r'[<>:"/\\|?*]', '', subject_name).strip()
    if clean_subj.endswith("-REV"):
        clean_subj = clean_subj[:-4].strip()

    wb = openpyxl.Workbook()
    
    # Colores Institucionales
    fill_navy = PatternFill(start_color="1A3A5C", end_color="1A3A5C", fill_type="solid")
    fill_amber = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
    fill_light = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_warn = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_title = Font(name="Arial", size=13, bold=True, color="1A3A5C")
    font_sub = Font(name="Arial", size=9, italic=True, color="475569")
    font_data = Font(name="Arial", size=9)
    font_bold = Font(name="Arial", size=9, bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    # -------------------------------------------------------------
    # HOJA 1: Matriz_Control_Documentos
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Matriz_Control_Documentos"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1.merge_cells("A1:I1")
    ws1["A1"] = "SISTEMA DE GESTIÓN DE LA CALIDAD EDUCATIVA • ISO 21001:2018 (CLÁUSULA 7.5)"
    ws1["A1"].font = Font(name="Arial", size=9, bold=True, color="B45309")
    ws1["A1"].alignment = Alignment(horizontal="center")
    
    ws1.merge_cells("A2:I2")
    ws1["A2"] = f"MATRIZ MAESTRA DE CONTROL DOCUMENTAL — ASIGNATURA: {clean_subj.upper()}"
    ws1["A2"].font = font_title
    ws1["A2"].alignment = Alignment(horizontal="center")

    ws1.merge_cells("A3:I3")
    ws1["A3"] = f"Registro Oficial de Versiones Vigentes y Trazabilidad de Evidencias Curriculares • Emisión: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws1["A3"].font = font_sub
    ws1["A3"].alignment = Alignment(horizontal="center")

    headers_ws1 = [
        "Código Documento",
        "Nombre Oficial del Documento",
        "Tipo de Documento",
        "Versión Vigente",
        "Fecha Modificación",
        "Responsable de Edición",
        "Origen del Cambio",
        "Hash SHA-256 (Integridad)",
        "Estado Aprobación"
    ]
    
    for col_num, h_text in enumerate(headers_ws1, 1):
        cell = ws1.cell(row=5, column=col_num, value=h_text)
        cell.fill = fill_navy
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws1.row_dimensions[5].height = 28

    # -------------------------------------------------------------
    # HOJA 2: Historial_Trazabilidad_ISO
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Historial_Trazabilidad_ISO")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:G1")
    ws2["A1"] = f"BITÁCORA INMUTABLE DE AUDITORÍA Y TRAZABILIDAD DOCUMENTAL — {clean_subj.upper()}"
    ws2["A1"].font = font_title
    ws2["A1"].alignment = Alignment(horizontal="center")

    ws2.merge_cells("A2:G2")
    ws2["A2"] = "Historial cronológico de cambios generados automáticamente por agentes de IA y por el docente titular."
    ws2["A2"].font = font_sub
    ws2["A2"].alignment = Alignment(horizontal="center")

    headers_ws2 = [
        "ID Evento",
        "Fecha y Hora",
        "Documento Afectado",
        "Transición Versión",
        "Autor / Agente Responsable",
        "Descripción del Cambio Realizado",
        "Justificación Pedagógica del Cambio"
    ]
    
    for col_num, h_text in enumerate(headers_ws2, 1):
        cell = ws2.cell(row=4, column=col_num, value=h_text)
        cell.fill = fill_navy
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws2.row_dimensions[4].height = 28

    # -------------------------------------------------------------
    # HOJA 3: Normativa_y_Advertencias
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Normativa_y_Advertencias")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:F1")
    ws3["A1"] = "DIRECTIVAS DE CONTROL DOCUMENTAL Y GESTIÓN DE CALIDAD EDUCATIVA"
    ws3["A1"].font = font_title
    ws3["A1"].alignment = Alignment(horizontal="center")

    # Banner de advertencia sobre edición manual
    ws3.merge_cells("A3:F5")
    warn_text = (
        "⚠️ AVISO OBLIGATORIO DE AUDITORÍA INSTITUCIONAL (ISO 21001 / MODELO DE ACREDITACIÓN):\n"
        "Si usted como docente o coordinador académico edita manualmente cualquier archivo Word (.docx) o Excel (.xlsx) "
        "fuera de la plataforma (por ejemplo, abriéndolo directamente en Microsoft Word, LibreOffice o Google Docs), ES OBLIGATORIO "
        "registrar el cambio en esta bitácora actualizando el número de versión (ej. 1.0 a 1.1), la fecha y la justificación pedagógica.\n"
        "Cualquier discrepancia entre el Hash/Fecha del archivo físico y esta matriz invalida la trazabilidad del portafolio en auditorías de acreditación."
    )
    ws3["A3"] = warn_text
    ws3["A3"].font = Font(name="Arial", size=9.5, bold=True, color="78350F")
    ws3["A3"].fill = fill_warn
    ws3["A3"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    normativa_rows = [
        ("Criterio de Versionado", "Definición Institucional", "Ejemplo Práctico"),
        ("Versión Mayor (X.0)", "Cambios estructurales en competencias globales, cambio de ponderaciones en el plan de evaluación o reemplazo de instrumentos sumativos.", "De 1.0 a 2.0 al rediseñar las 4 fases curriculares."),
        ("Versión Menor (X.y)", "Ajustes didácticos de redacción, corrección de erratas, inclusión de nuevas pautas DUA para un estudiante específico o recalibración de rúbrica.", "De 1.0 a 1.1 al añadir adaptación DUA de tiempo extendido."),
        ("Pipeline IA Autónomo", "Documentos compilados o procesados directamente mediante los agentes (e1, e2, e3, e4). El sistema calcula el Hash SHA-256 automáticamente.", "Generación de BASE_INTEGRADA.xlsx tras limpieza de datos."),
        ("Edición Manual Docente", "Intervenciones realizadas por el docente titular que enriquecen o contextualizan los productos generados por la IA.", "Revisión y personalización del Plan Modular en Word.")
    ]

    for r_idx, row_vals in enumerate(normativa_rows, start=7):
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws3.cell(row=r_idx, column=c_idx, value=val)
            if r_idx == 7:
                cell.fill = fill_amber
                cell.font = font_header
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.fill = fill_light if r_idx % 2 == 0 else PatternFill(fill_type=None)
                cell.font = font_data
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws3.row_dimensions[r_idx].height = 32

    # Autoajustar anchos de columna en las tres hojas
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len and "\n" not in val and cell.coordinate not in ["A1", "A2", "A3"]:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    ws1.column_dimensions["B"].width = 38
    ws1.column_dimensions["A"].width = 24
    ws1.column_dimensions["H"].width = 24
    ws2.column_dimensions["C"].width = 35
    ws2.column_dimensions["F"].width = 40
    ws2.column_dimensions["G"].width = 40
    ws3.column_dimensions["A"].width = 25
    ws3.column_dimensions["B"].width = 50
    ws3.column_dimensions["C"].width = 40

    wb.save(wb_path)
    return wb_path

def record_document_version(
    subject_name: str,
    doc_name: str,
    doc_code: str,
    doc_type: str,
    new_version: str,
    author_or_agent: str,
    change_description: str,
    justification: str = "Aseguramiento de la calidad curricular y evidencias de aprendizaje.",
    file_path: Optional[str] = None,
    approval_status: str = "Aprobado Institucional"
) -> Dict[str, Any]:
    """
    Registra o actualiza un documento en la Matriz de Control y añade un evento al Historial ISO.
    """
    wb_path = get_version_control_path(subject_name)
    if not os.path.exists(wb_path):
        init_version_control_workbook(subject_name)
        
    wb = openpyxl.load_workbook(wb_path)
    ws_matrix = wb["Matriz_Control_Documentos"]
    ws_history = wb["Historial_Trazabilidad_ISO"]
    
    now_str = datetime.now().strftime("%d/%m/%Y %H:%M")
    sha_hash = calculate_sha256(file_path) if file_path and os.path.exists(file_path) else "REGISTRADO_EN_DISCO"
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    font_data = Font(name="Arial", size=9)
    font_bold = Font(name="Arial", size=9, bold=True)
    
    # 1. Actualizar o insertar en Matriz_Control_Documentos
    row_found = None
    old_version = "0.0"
    for r in range(6, ws_matrix.max_row + 1):
        c_code = ws_matrix.cell(row=r, column=1).value
        c_name = ws_matrix.cell(row=r, column=2).value
        if c_code == doc_code or c_name == doc_name:
            row_found = r
            old_version = str(ws_matrix.cell(row=r, column=4).value or "1.0")
            break
            
    is_agent = ("Agente" in author_or_agent or "Pipeline" in author_or_agent or "e1" in author_or_agent or "e2" in author_or_agent or "e3" in author_or_agent or "e4" in author_or_agent)
    origin = "Pipeline IA Autónomo" if is_agent else "Edición Manual Docente"
    
    target_row = row_found if row_found else ws_matrix.max_row + 1
    row_data = [
        doc_code,
        doc_name,
        doc_type,
        new_version,
        now_str,
        author_or_agent,
        origin,
        sha_hash,
        approval_status
    ]
    
    for c_idx, val in enumerate(row_data, start=1):
        cell = ws_matrix.cell(row=target_row, column=c_idx, value=val)
        cell.font = font_bold if c_idx in [1, 4] else font_data
        cell.border = thin_border
        if c_idx in [1, 4, 5, 7, 9]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
    ws_matrix.row_dimensions[target_row].height = 22

    # 2. Agregar evento inmutable al Historial_Trazabilidad_ISO
    hist_row = ws_history.max_row + 1
    event_id = f"EVT-{hist_row - 4:03d}"
    
    hist_data = [
        event_id,
        now_str,
        doc_name,
        f"{old_version} -> {new_version}",
        author_or_agent,
        change_description,
        justification
    ]
    
    for c_idx, val in enumerate(hist_data, start=1):
        cell = ws_history.cell(row=hist_row, column=c_idx, value=val)
        cell.font = font_bold if c_idx in [1, 4] else font_data
        cell.border = thin_border
        if c_idx in [1, 2, 4]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
    ws_history.row_dimensions[hist_row].height = 24

    wb.save(wb_path)
    
    return {
        "status": "success",
        "subject": subject_name,
        "document": doc_name,
        "version": new_version,
        "event_id": event_id,
        "workbook_path": wb_path,
        "hash": sha_hash
    }

def get_subject_version_matrix(subject_name: str) -> List[Dict[str, Any]]:
    """Lee y devuelve los documentos vigentes registrados en la matriz de control."""
    wb_path = get_version_control_path(subject_name)
    if not os.path.exists(wb_path):
        init_version_control_workbook(subject_name)
        
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    ws = wb["Matriz_Control_Documentos"]
    
    results = []
    for r in range(6, ws.max_row + 1):
        code = ws.cell(row=r, column=1).value
        name = ws.cell(row=r, column=2).value
        if not code and not name:
            continue
        results.append({
            "codigo": str(code or ""),
            "nombre": str(name or ""),
            "tipo": str(ws.cell(row=r, column=3).value or ""),
            "version": str(ws.cell(row=r, column=4).value or "1.0"),
            "fecha": str(ws.cell(row=r, column=5).value or ""),
            "responsable": str(ws.cell(row=r, column=6).value or ""),
            "origen": str(ws.cell(row=r, column=7).value or ""),
            "hash": str(ws.cell(row=r, column=8).value or ""),
            "estado": str(ws.cell(row=r, column=9).value or "Vigente")
        })
    return results
